from __future__ import annotations

from pathlib import Path
from datetime import date
import csv

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
ASSET = ROOT / "outputs" / "technical_report_assets"
OUT = ROOT / "outputs" / "Bao_cao_tong_hop_de_tai_loai_bo_nhan_Xquang.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
DARK = "1F2937"
GRAY = "5B6573"
LIGHT_GRAY = "E9EDF2"
GREEN = "DFF1E6"
ORANGE = "FFF0D7"
RED = "FDE3E3"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float):
    cell.width = Cm(width_cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_cm / 2.54 * 1440)))
    tcW.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trPr.append(cant_split)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    rPr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_text(doc: Document, text: str, bold_prefix: str | None = None, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc: Document, text: str, level: int = 0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.7 + 0.45 * level)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text)
    return p


_number_counter = 0


def add_number(doc: Document, text: str):
    global _number_counter
    _number_counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(f"{_number_counter}. ").bold = True
    p.add_run(text)
    return p


def add_callout(doc: Document, title: str, body: str, fill=PALE_BLUE):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    prevent_row_split(t.rows[0])
    cell = t.cell(0, 0)
    set_cell_width(cell, 16.2)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 140, 180, 140, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p2 = cell.add_paragraph(body)
    if "\n" in body:
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p2.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9.2)
    else:
        p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, image: Path, caption: str, width_cm=16.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(image), width=Cm(width_cm))
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = False
    return cap


def add_table(doc: Document, headers, rows, widths=None, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, h in enumerate(headers):
        c = hdr.cells[j]
        c.text = str(h)
        set_cell_shading(c, BLUE)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(font_size)
        if widths:
            set_cell_width(c, widths[j])
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            c = cells[j]
            c.text = str(value)
            set_cell_margins(c)
            if widths:
                set_cell_width(c, widths[j])
            if i % 2 == 1:
                set_cell_shading(c, "F7F9FB")
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_equation(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)


def add_heading(doc: Document, text: str, level=1):
    global _number_counter
    _number_counter = 0
    p = doc.add_heading(text, level=level)
    keep_with_next(p)
    return p


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color, before, after in [
        ("Title", 24, BLUE, 0, 10),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 11, 5),
        ("Heading 3", 11.5, DARK, 8, 4),
    ]:
        st = doc.styles[name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    cap = doc.styles["Caption"]
    cap.font.name = "Times New Roman"
    cap.font.size = Pt(9.5)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string(GRAY)
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(8)

    for sname in ("List Bullet", "List Bullet 2", "List Number"):
        st = doc.styles[sname]
        st.font.name = "Times New Roman"
        st.font.size = Pt(11)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.paragraph_format.space_after = Pt(3)

    header = sec.header.paragraphs[0]
    header.text = "BÁO CÁO KỸ THUẬT • LOẠI BỎ NHÃN PHI GIẢI PHẪU TRÊN X-QUANG BÀN TAY"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(sec.footer.paragraphs[0])

    # Widow/orphan control for body styles.
    for st in doc.styles:
        if getattr(st, "type", None) == WD_STYLE_TYPE.PARAGRAPH:
            pPr = st.element.get_or_add_pPr()
            widow = pPr.find(qn("w:widowControl"))
            if widow is None:
                widow = OxmlElement("w:widowControl")
                pPr.append(widow)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    r = p.add_run("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("BÁO CÁO TỔNG HỢP KỸ THUẬT")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    doc.add_paragraph()
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = line.add_run("LOẠI BỎ NHÃN PHI GIẢI PHẪU\nTRÊN ẢNH X-QUANG BÀN TAY")
    rr.bold = True
    rr.font.size = Pt(24)
    rr.font.color.rgb = RGBColor.from_string(DARK)
    line.paragraph_format.space_after = Pt(14)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Phương pháp chỉnh sửa cục bộ bảo toàn giải phẫu\nvà đánh giá tác động lên ước lượng tuổi xương")
    sr.italic = True
    sr.font.size = Pt(14)
    sr.font.color.rgb = RGBColor.from_string(BLUE)

    if (ASSET / "pipeline_artifact_only.png").exists():
        pimg = doc.add_paragraph()
        pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pimg.add_run().add_picture(str(ASSET / "pipeline_artifact_only.png"), width=Cm(15.4))

    info = doc.add_table(rows=4, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    entries = [
        ("Nhóm thực hiện", "Nguyễn Nhật Toàn – Đinh Tấn Phương – Lê Trần Phú"),
        ("Tập đánh giá", "200 ảnh RSNA, Case_ID 4360–4559"),
        ("Phạm vi", "Tiền xử lý ảnh y khoa và đánh giá tuổi xương"),
        ("Ngày tổng hợp", "29/07/2026"),
    ]
    for i, (k, v) in enumerate(entries):
        set_cell_width(info.cell(i, 0), 4.0)
        set_cell_width(info.cell(i, 1), 11.3)
        set_cell_shading(info.cell(i, 0), LIGHT_BLUE)
        set_cell_margins(info.cell(i, 0), 100, 140, 100, 140)
        set_cell_margins(info.cell(i, 1), 100, 140, 100, 140)
        info.cell(i, 0).text = k
        info.cell(i, 1).text = v
        info.cell(i, 0).paragraphs[0].runs[0].bold = True
    pfoot = doc.add_paragraph()
    pfoot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pfoot.paragraph_format.space_before = Pt(18)
    rr = pfoot.add_run("TP. Hồ Chí Minh, 2026")
    rr.font.size = Pt(11)
    rr.font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_page_break()


def load_summary():
    with open(ASSET / "summary_three_way.csv", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    order_protocol = ["fullframe", "histmatch", "fixed_crop", "fixed_crop_histmatch"]
    order_group = ["original", "ours", "author_mean"]
    d = {(r["Protocol"], r["Group"]): r for r in rows}
    return [d[(p, g)] for p in order_protocol for g in order_group]


def load_paired():
    with open(ASSET / "paired_three_way.csv", encoding="utf-8-sig", newline="") as f:
        return {(r["Protocol"], r["Group"]): r for r in csv.DictReader(f)}


def protocol_label(p):
    return {
        "fullframe": "Toàn khung",
        "histmatch": "Toàn khung + khớp histogram",
        "fixed_crop": "Crop cố định",
        "fixed_crop_histmatch": "Crop cố định + khớp histogram",
    }[p]


def group_label(g):
    return {"original": "Ảnh gốc", "ours": "Ảnh làm sạch", "author_mean": "Ảnh tác giả (TB 3 lần)"}[g]


def build():
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "Tóm tắt điều hành", 1)
    add_text(
        doc,
        "Báo cáo này ghi lại toàn bộ quá trình chuyển từ phương án “tách toàn bộ bàn tay rồi ghép sang nền mới” sang phương án “chỉ chỉnh sửa vùng dị vật, khóa cứng vùng giải phẫu”. "
        "Thay đổi này xuất phát trực tiếp từ các lỗi mất da, thủng bàn tay, biên răng cưa, nền giả và loại bỏ ảnh khi phân đoạn thất bại. "
        "Giải pháp cuối cùng dùng mask dị vật được duyệt thủ công, mask bảo vệ giải phẫu, mô hình nền cục bộ bậc một kèm grain và cơ chế ghép bảo toàn pixel.",
    )
    add_callout(
        doc,
        "Kết quả chính",
        "Đã xử lý đủ 200/200 ảnh; 185 ảnh dùng mask thủ công và 15 ảnh làm sạch tự động. "
        "Kiểm thử bất biến đạt 200/200: không có pixel nào thay đổi trong vùng giải phẫu được bảo vệ và không có pixel nào thay đổi ngoài mask dị vật. "
        "Trên mô hình ianpan/bone-age, MAE toàn khung của ảnh gốc là 4,42 tháng, ảnh làm sạch là 4,32 tháng; độ trôi dự đoán trung bình chỉ 0,76 tháng. "
        "Ngược lại, trung bình ba ảnh inpainting của tác giả có MAE 43,40 tháng và độ trôi 42,14 tháng với cùng checkpoint đánh giá.",
        GREEN,
    )
    add_text(
        doc,
        "Kết luận cốt lõi: với mục tiêu loại bỏ nhãn phi giải phẫu để đánh giá tuổi xương, ưu tiên số một không phải là tạo ra một ảnh “đẹp”, mà là giới hạn miền được phép sửa và chứng minh bằng kiểm thử rằng cấu trúc giải phẫu không bị thay đổi. "
        "Phương pháp cục bộ của đề tài thỏa mãn tinh thần này tốt hơn inpainting toàn ảnh hoặc ghép toàn bộ bàn tay sang nền mới.",
    )

    add_heading(doc, "Mục lục nội dung", 1)
    toc_items = [
        "1. Bối cảnh, bài toán và nguồn gốc ý tưởng",
        "2. Phân tích các hướng tiếp cận và quyết định thiết kế",
        "3. Pipeline cuối cùng: chỉnh sửa cục bộ bảo toàn giải phẫu",
        "4. Các sự cố trong quá trình phát triển và cách khắc phục",
        "5. Kết quả xử lý 200 ảnh và kiểm soát chất lượng",
        "6. Mô hình tuổi xương dùng để so sánh",
        "7. Thiết kế thực nghiệm và ý nghĩa các phép đo",
        "8. Kết quả so sánh định lượng và phân tích biểu đồ",
        "9. Cách cài đặt, chạy và tái lập kết quả",
        "10. Hạn chế, khuyến nghị và kết luận cuối cùng",
        "Tài liệu tham khảo và phụ lục đường dẫn",
    ]
    for item in toc_items:
        add_bullet(doc, item)
    doc.add_page_break()

    add_heading(doc, "1. Bối cảnh, bài toán và nguồn gốc ý tưởng", 1)
    add_heading(doc, "1.1. Bài toán nghiên cứu", 2)
    add_text(
        doc,
        "Ảnh X-quang bàn tay dùng cho ước lượng tuổi xương thường chứa ký tự chỉ bên trái/phải, nhãn bệnh viện, băng marker, viền cassette, thiết bị cố định hoặc vật thể ngoài giải phẫu. "
        "Những thành phần này có thể trở thành tín hiệu nhiễu đối với mô hình học sâu. Mục tiêu của đề tài là tạo một ảnh tương ứng cho từng ảnh gốc trong bộ 200 ca của tác giả, loại bỏ nhãn phi giải phẫu nhưng không thay đổi xương, khe sụn tăng trưởng, mô mềm hoặc tư thế bàn tay.",
    )
    add_text(
        doc,
        "Yêu cầu khoa học quan trọng nhất là tính đối ứng theo ca: cùng Case_ID, cùng giới tính, cùng ground truth tuổi, chỉ khác thao tác làm sạch. "
        "Nếu ảnh bị loại khỏi pipeline hoặc giải phẫu bị tái tạo lại, phép so sánh tuổi xương sẽ không còn là một thí nghiệm ghép cặp công bằng.",
    )

    add_heading(doc, "1.2. Nguồn gốc ý tưởng từ bài báo chính", 2)
    add_text(
        doc,
        "Bài báo “Evaluating the Clinical Impact of Generative Inpainting on Bone Age Estimation” của Matsuoka và cộng sự [1] dùng GPT-image-1 để tạo ba biến thể cho mỗi ảnh trong 200 ảnh RSNA. "
        "Ảnh được đổi kích thước 1024 × 1024, mask inpainting là toàn màu trắng—nghĩa là mô hình được phép sửa toàn bộ ảnh—và prompt yêu cầu loại bỏ artefact đồng thời giữ giải phẫu.",
    )
    add_callout(
        doc,
        "Tại sao bài báo tạo ra động lực cho đề tài?",
        "Bài báo cho thấy một prompt bảo toàn giải phẫu không đủ tạo ràng buộc pixel. Mô hình sinh có thể làm trưởng thành xương, làm mờ hoặc hợp nhất sụn tăng trưởng, thay đổi hình dạng xương và tạo biến thiên giữa các lần sinh. "
        "MAE của mô hình đánh giá trong bài báo tăng từ 6,26 lên 30,11 tháng; độ lệch chuẩn dự đoán giữa ba lần sinh trung bình 8,77 ± 5,81 tháng.",
        ORANGE,
    )
    add_text(
        doc,
        "Do đó, ý tưởng trung tâm của đề tài không phải là huấn luyện một mô hình sinh mạnh hơn, mà là thay đổi phạm vi chỉnh sửa: từ “mô hình có thể sửa mọi pixel” sang “chỉ một mask dị vật nhỏ được phép thay đổi; vùng giải phẫu bị khóa cứng”. "
        "Đây là nguồn gốc trực tiếp của tư tưởng artifact-only inpainting.",
    )

    add_heading(doc, "1.3. Nguồn dữ liệu và độ tin cậy", 2)
    rows = [
        ["Ảnh gốc + tuổi/giới tính", "RSNA test, 200 ca 4360–4559", "CSV data/rsna_test.csv từ fork sRassmann/bone-age; khớp đúng 200 ID; dùng làm ground truth."],
        ["Ảnh inpainting tác giả", "600 PNG = 200 ca × 3 lần sinh", "Tên file, kích thước từng file và tổng dung lượng khớp Kaggle chính thức của Felipe Matsuoka."],
        ["Mã bài báo", "Repo chính thức Felipe Matsuoka", "Mô tả quy trình tạo ảnh, không phải nguồn duy nhất lưu toàn bộ ảnh."],
        ["Repo nhóm", "nhattoan235/gpt-image-bone-age-synthesis", "Repo làm việc của nhóm, dẫn chiếu dữ liệu; không nên dùng như bằng chứng độc lập về nguồn gốc 600 ảnh."],
    ]
    add_table(doc, ["Thành phần", "Quy mô", "Đánh giá độ tin cậy"], rows, [3.4, 4.3, 8.5], 8.6)
    add_text(
        doc,
        "Kết luận về provenance: bộ 600 ảnh cục bộ có độ tin cậy cao để dùng làm ảnh của tác giả vì cấu trúc tên, số lượng, kích thước từng file và tổng byte đều khớp danh mục Kaggle chính thức. "
        "Tuy nhiên Kaggle không công bố checksum SHA-256 cho từng file, vì vậy báo cáo tránh dùng từ “đồng nhất mật mã”; mức kết luận phù hợp là “khớp metadata đầy đủ”.",
    )

    add_heading(doc, "2. Phân tích các hướng tiếp cận và quyết định thiết kế", 1)
    add_heading(doc, "2.1. Hướng của tác giả: inpainting toàn ảnh", 2)
    add_text(
        doc,
        "Ưu điểm của inpainting toàn ảnh là thao tác đơn giản và có khả năng tạo nền mượt. Nhược điểm quyết định là không có cơ chế bảo đảm giải phẫu. "
        "Mask trắng toàn ảnh biến yêu cầu “chỉ xóa nhãn” thành một bài toán tái tổng hợp ảnh X-quang. Mọi thay đổi trong tuổi xương dự đoán sau đó có thể đến từ việc xóa nhãn, thay đổi texture, thay đổi xương hoặc tất cả các yếu tố trên.",
    )

    add_heading(doc, "2.2. Hướng ban đầu của nhóm: phân đoạn và ghép nền", 2)
    add_text(
        doc,
        "Phiên bản đầu tiên phân ngưỡng bàn tay, sửa mask bằng morphology, giữ nguyên pixel bên trong mask rồi ghép toàn bộ bàn tay sang nền Gaussian hoặc nền lấy từ ngân hàng ảnh sạch. "
        "Tư tưởng này đúng ở chỗ cố tránh sinh lại xương, nhưng phạm vi can thiệp vẫn quá lớn: toàn bộ nền bị thay mới và chất lượng phụ thuộc tuyệt đối vào biên mask bàn tay.",
    )
    add_table(
        doc,
        ["Chỉ số pipeline cũ", "Kết quả", "Hệ quả"],
        [
            ["Số ảnh đầu vào", "200", "Mục tiêu cần đủ 200 đầu ra đối ứng."],
            ["SUCCESS", "177", "Chỉ 88,5% ảnh được sinh."],
            ["SEGMENTATION_FAILED", "23", "11,5% bị loại; làm sai thiết kế ghép cặp."],
            ["Nhãn chồng vùng tay", "34", "Khó quyết định giữ nhãn hay làm mất giải phẫu."],
            ["Thời gian trung bình A/B", "113,8 / 328,1 ms", "Tốc độ tốt nhưng không bù được sai số hình thái."],
        ],
        [5.0, 3.2, 8.0],
        8.8,
    )
    add_figure(
        doc,
        ASSET / "trouble_and_fix_montage.png",
        "Hình 1. Tổng hợp lỗi của pipeline cũ và kết quả sau khi chuyển sang chỉnh sửa cục bộ. Ảnh ghép nền cho thấy đường viền răng cưa, nền phẳng và vùng mô mềm bị cắt; panel cuối cho thấy chỉ vùng nhãn được thay đổi.",
        16.2,
    )

    add_heading(doc, "2.3. So sánh với hướng Stable Diffusion + ControlNet của bạn cùng nhóm", 2)
    add_text(
        doc,
        "Hướng được đề xuất bởi bạn cùng nhóm gồm SAM phân đoạn bàn tay, phát hiện nhãn ngoài bàn tay, tạo nền thô, Stable Diffusion Inpainting với ControlNet Canny và LoRA X-quang, sau đó ghép lại pixel gốc ngoài mask. "
        "Hai hướng giống nhau ở ba ý: xác định vùng cần xóa, dựng nền và ghép lại pixel gốc ngoài mask. Điểm khác cốt lõi là đề tài hiện tại không dùng mô hình sinh để lấp vùng mask; nó dùng mô hình nền cục bộ có thể giải thích và kiểm soát.",
    )
    add_table(
        doc,
        ["Tiêu chí", "Stable Diffusion + ControlNet", "Artifact-only của đề tài"],
        [
            ["Miền sửa", "Mask nhãn", "Mask nhãn, tiếp tục bị chặn bởi mask bảo vệ giải phẫu"],
            ["Cơ chế lấp", "Mô hình sinh có prior học được", "Mặt phẳng gradient cục bộ + residual/grain"],
            ["Rủi ro", "Hallucination, phụ thuộc LoRA/seed", "Nền đơn giản ở mask rất lớn; không suy diễn giải phẫu"],
            ["Khả năng chứng minh", "Cần kiểm tra hậu nghiệm", "Bất biến pixel được kiểm tra trực tiếp"],
            ["Phù hợp đồ án", "Mạnh về sinh ảnh", "Mạnh về kiểm soát tiền xử lý y khoa"],
        ],
        [3.2, 6.4, 6.6],
        8.4,
    )
    add_callout(
        doc,
        "Lý do chọn artifact-only",
        "Mục tiêu nghiên cứu là đánh giá ảnh hưởng của việc loại bỏ nhãn lên tuổi xương, không phải chứng minh khả năng sinh ảnh. "
        "Vì vậy phương pháp tốt nhất là phương pháp thay đổi ít pixel nhất, có ràng buộc không sửa giải phẫu và có thể kiểm toán từng ca.",
        GREEN,
    )

    add_heading(doc, "3. Pipeline cuối cùng: chỉnh sửa cục bộ bảo toàn giải phẫu", 1)
    add_figure(
        doc,
        ASSET / "pipeline_artifact_only.png",
        "Hình 2. Pipeline artifact-only. Nhánh mask bảo vệ hoạt động như một hàng rào không cho phép vùng giải phẫu bị chỉnh sửa; nhánh mask dị vật xác định miền duy nhất được phép tái tạo.",
        16.2,
    )
    add_heading(doc, "3.1. Tư tưởng bảo toàn tối thiểu", 2)
    add_text(
        doc,
        "Gọi ảnh gốc là I, mask dị vật là Mₐ và mask bảo vệ giải phẫu là Mₚ. Mask chỉnh sửa thực tế được định nghĩa M = Mₐ ∩ ¬Mₚ. "
        "Ảnh nền ước lượng trong mask là B. Ảnh đầu ra O được lắp ráp theo quy tắc: ngoài M giữ nguyên I; trong M mới dùng B và feather hẹp ở phía trong. "
        "Do đó O(x)=I(x) với mọi pixel x∉M, và đặc biệt O(x)=I(x) với mọi x∈Mₚ.",
    )
    add_equation(doc, "M = Mₐ ∩ ¬Mₚ ;    O(x) = I(x), ∀x ∉ M")
    add_text(
        doc,
        "Khác với lời hứa bằng prompt, đây là ràng buộc xác định trong mã. Hai invariant được ghi log cho từng ảnh: changed_inside_protected = 0 và changed_outside_artifact = 0.",
    )

    add_heading(doc, "3.2. Tạo mask bảo vệ giải phẫu", 2)
    add_text(
        doc,
        "Mask bảo vệ được tạo bảo thủ: ưu tiên bao phủ dư da, mô mềm, xương bàn tay và cẳng tay hơn là bám sát biên. "
        "Quy trình dùng tương phản toàn cục/cục bộ, thành phần liên thông chính, đóng hình thái để nối vùng đứt và dilation để tạo biên an toàn. "
        "Mask này không dùng để cắt bàn tay ra khỏi ảnh; nó chỉ định nghĩa vùng cấm sửa.",
    )
    add_bullet(doc, "Closing được thực hiện trước opening để nối các khe nhỏ mà không xóa ngón tay mảnh.")
    add_bullet(doc, "Giữ thành phần liên thông chính nhưng không ép phải là một silhouette hoàn hảo.")
    add_bullet(doc, "Dilation tạo safety margin quanh mô mềm; chấp nhận sót nhãn sát da còn hơn xóa nhầm giải phẫu.")

    add_heading(doc, "3.3. Phát hiện và duyệt mask dị vật", 2)
    add_text(
        doc,
        "Bộ phát hiện tự động tìm các nét sáng/tối bất thường theo tương phản cục bộ, gom thành phần gần nhau và áp dụng quy tắc hình học. "
        "Tuy nhiên dị vật X-quang rất đa dạng: chữ, băng marker, vòng kim loại, ống/kim, vật cố định và nhãn có thể chồng lên bàn tay. "
        "Vì vậy bước duyệt thủ công được thiết kế như một phần của phương pháp, không phải thao tác sửa ảnh bằng tay.",
    )
    add_callout(
        doc,
        "Ý nghĩa của việc tô đỏ",
        "Người dùng chỉ chú thích mask dị vật—tức là chỉ ra pixel nào thuộc nhãn/thiết bị phi giải phẫu. Không tô bàn tay và không vẽ giá trị pixel đầu ra. "
        "Đây là annotation phục vụ tiền xử lý có kiểm soát, hoàn toàn phù hợp tinh thần đồ án.",
        PALE_BLUE,
    )
    add_text(
        doc,
        "Với dị vật nằm trên hình chiếu bàn tay (ví dụ kim/ốc), hệ thống không được phép “đoán” phần xương bị che. "
        "Mask được vẽ đúng theo dị vật, sau đó bị cắt bởi mask bảo vệ. Phần chồng giải phẫu được giữ nguyên hoặc đánh dấu ca đặc biệt để loại khỏi phân tích phụ; không dùng inpainting sinh để tái tạo xương không quan sát được.",
    )

    add_heading(doc, "3.4. Tái tạo nền bằng gradient cục bộ và grain", 2)
    add_text(
        doc,
        "Đối với mỗi thành phần mask, hệ thống lấy một vành mẫu bao quanh vùng cần xóa. Trên các pixel hợp lệ của vành, cường độ được xấp xỉ bằng mặt phẳng bậc một theo tọa độ:",
    )
    add_equation(doc, "B(x,y) = β₀ + β₁x + β₂y")
    add_text(
        doc,
        "Ba hệ số β được ước lượng từ nền lân cận. β₁ và β₂ mô tả gradient chiếu sáng theo chiều ngang và dọc; β₀ là mức nền. "
        "Sau đó residual của vành—phần sai khác giữa pixel thật và mặt phẳng—được lấy mẫu/lọc để bổ sung grain X-quang. "
        "Nhờ đó vùng lấp không phẳng tuyệt đối và vẫn có texture thống kê gần vùng xung quanh.",
    )
    add_text(
        doc,
        "Cuối cùng chỉ một dải feather hẹp nằm bên trong mask được trộn với ảnh gốc. Dùng feather hướng vào trong rất quan trọng: nếu blur alpha ra ngoài mask, pixel hợp lệ ngoài dị vật cũng bị thay đổi và kiểm thử bất biến sẽ thất bại.",
    )

    add_heading(doc, "3.5. Ghép đầu ra và QC theo pixel", 2)
    add_number(doc, "Đọc ảnh theo thang xám và giữ nguyên kích thước gốc.")
    add_number(doc, "Tạo mask bảo vệ; nạp mask dị vật tự động hoặc mask thủ công đã duyệt.")
    add_number(doc, "Tính mask chỉnh sửa M = Mₐ ∩ ¬Mₚ.")
    add_number(doc, "Ước lượng nền cục bộ cho từng thành phần và bổ sung grain.")
    add_number(doc, "Seam-match và feather hẹp phía trong mask.")
    add_number(doc, "Ép sao chép lại pixel gốc ngoài M và toàn bộ Mₚ.")
    add_number(doc, "Xuất ảnh, mask, protected mask, review panel và dòng QC trong CSV.")

    add_heading(doc, "4. Các sự cố trong quá trình phát triển và cách khắc phục", 1)
    add_heading(doc, "4.1. Mất ngón tay và thủng mask do morphology", 2)
    add_text(
        doc,
        "Trong pipeline cũ, opening (erosion rồi dilation) được áp dụng quá sớm. Với các ngón tay mảnh hoặc có cường độ thấp, erosion làm đứt kết nối; dilation sau đó không thể khôi phục phần đã mất. "
        "Các lỗ nhỏ trong mô mềm cũng trở thành lỗ thật trong silhouette, dẫn tới nền xuyên vào bàn tay.",
    )
    add_text(
        doc,
        "Cách sửa trung gian là closing trước opening để nối khe và lấp lỗ. Điều này cải thiện mask nhưng không giải quyết được vấn đề nền tảng: một lỗi biên chỉ 2–3 pixel vẫn làm mất da khi ghép toàn bộ bàn tay. "
        "Giải pháp cuối cùng là ngừng dùng mask bàn tay làm alpha compositing; mask bảo vệ chỉ là hàng rào cấm sửa.",
    )

    add_heading(doc, "4.2. 23 ảnh không sinh được", 2)
    add_text(
        doc,
        "Otsu/threshold đôi khi chọn cassette, viền ảnh hoặc một vùng nền sáng thay vì bàn tay; các bộ lọc hình học sau đó loại ảnh vì tỷ lệ diện tích, vị trí hoặc số thành phần không đạt. "
        "Kết quả là chỉ 177/200 ảnh có đầu ra, gây selection bias và không thể so sánh đúng với 200 ảnh tác giả.",
    )
    add_text(
        doc,
        "Cách giải quyết: đổi quan niệm từ “mask bàn tay phải hoàn hảo mới được sinh ảnh” sang “mọi ảnh đều có đầu ra; ca không chắc chắn chuyển sang review”. "
        "Mask dị vật thủ công là fallback an toàn. Pipeline cuối cùng không loại bất kỳ ca nào.",
    )

    add_heading(doc, "4.3. Nền phẳng, halo, biên răng cưa và mất da", 2)
    add_text(
        doc,
        "Ghép toàn bộ bàn tay sang nền Gaussian làm mất gradient chiếu xạ vốn có. Chênh lệch mean/variance tạo halo; alpha lấy từ mask nhị phân làm đường viền răng cưa; erosion quá mức cắt mô mềm. "
        "Ngay cả khi xương bên trong mask giữ pixel gốc, mô hình tuổi xương vẫn có thể phản ứng với crop, silhouette và phân bố cường độ mới.",
    )
    add_text(
        doc,
        "Cách giải quyết là không thay nền toàn khung. Chỉ vùng nhãn nhỏ được lấp từ vành cục bộ, nên gradient nền tổng thể, khung cassette và toàn bộ mô mềm ngoài mask vẫn nguyên bản.",
    )

    add_heading(doc, "4.4. Các thử nghiệm tái tạo nền không đạt", 2)
    add_figure(
        doc,
        ASSET / "reconstruction_trials.png",
        "Hình 3. Các thử nghiệm tái tạo nền. Telea dễ tạo vệt kéo; feather rộng để lại bóng chữ mờ; plane không khống chế có thể ngoại suy quá tối/sáng; phiên bản cuối dùng ring sạch, kẹp robust, grain và feather hẹp.",
        16.2,
    )
    add_table(
        doc,
        ["Thử nghiệm", "Vấn đề quan sát", "Nguyên nhân", "Điều chỉnh"],
        [
            ["Telea", "Nhòe hoặc kéo nét nhãn vào vùng lấp", "Lan truyền thông tin biên theo khoảng cách; biên vẫn chứa marker", "Không dùng cho vùng nhãn lớn/đậm"],
            ["Feather rộng", "Ghost chữ còn nhìn thấy", "Alpha trộn lại quá nhiều pixel nhãn gốc", "Chỉ feather dải hẹp phía trong"],
            ["Plane thuần", "Mảng quá tối/sáng gần biên cassette", "Vành mẫu trộn nhiều chế độ nền; ngoại suy không giới hạn", "Loại outlier, kẹp theo quantile cục bộ"],
            ["Không grain", "Mảng lấp phẳng và dễ nhận biết", "Mặt phẳng bậc một không chứa nhiễu detector", "Bổ sung residual/grain từ vành"],
            ["Bản cuối", "Seam nhỏ, không sửa ngoài mask", "Ring sạch + plane + grain + seam-match", "Dùng cho bộ 200 ảnh"],
        ],
        [2.5, 4.1, 5.0, 4.6],
        8.1,
    )

    add_heading(doc, "4.5. Lỗi giao diện annotation sau khi khởi động lại máy", 2)
    add_text(
        doc,
        "Giao diện mở được HTML nhưng hiển thị “Failed to fetch” hoặc 127.0.0.1 refused to connect. Nguyên nhân là file index.html chỉ là frontend; API Python không tự chạy sau khi Windows khởi động lại. "
        "Khi mở bằng đường dẫn file://, trình duyệt cũng không thể gọi đúng endpoint và bị giới hạn bởi origin.",
    )
    add_text(
        doc,
        "Cách giải quyết là khởi động annotation_server.py, giữ cửa sổ terminal chạy và truy cập đúng http://127.0.0.1:8765/. "
        "Frontend được bổ sung trạng thái retry và thông báo rõ máy chủ chưa phản hồi. Đây là lỗi vận hành, không làm mất các mask đã lưu.",
    )

    add_heading(doc, "5. Kết quả xử lý 200 ảnh và kiểm soát chất lượng", 1)
    add_figure(
        doc,
        ASSET / "qc_progress.png",
        "Hình 4. Tiến trình từ pipeline cũ sang bản cuối. Cột biểu diễn số ảnh đầu ra; các nhãn trên cột nêu trạng thái QC. Mục tiêu không chỉ là đạt 200 ảnh mà còn là đạt 200 kiểm thử bảo toàn pixel.",
        15.8,
    )
    add_table(
        doc,
        ["Chỉ số", "Giá trị", "Ý nghĩa"],
        [
            ["Ảnh đầu ra", "200/200", "Không còn loại ảnh; giữ thiết kế ghép cặp."],
            ["Mask thủ công", "185", "Các ca có dị vật được người dùng duyệt/vẽ mask."],
            ["Làm sạch tự động", "15", "Không cần sửa mask thủ công."],
            ["Diện tích mask dị vật TB", "3,123%", "Phạm vi can thiệp nhỏ so với toàn ảnh."],
            ["Pixel thực sự thay đổi TB", "2,676%", "Nhỏ hơn diện tích mask vì có vùng được bảo vệ/feather."],
            ["Diện tích bảo vệ TB", "40,932%", "Bao phủ bàn tay/cẳng tay theo hướng bảo thủ."],
            ["changed_inside_protected", "0 ở 200/200", "Không sửa pixel trong vùng giải phẫu bảo vệ."],
            ["changed_outside_artifact", "0 ở 200/200", "Không sửa pixel ngoài vùng dị vật."],
        ],
        [5.0, 3.2, 8.0],
        8.7,
    )
    add_callout(
        doc,
        "Diễn giải đúng về “bảo toàn giải phẫu”",
        "QC pixel chứng minh vùng được định nghĩa là giải phẫu bảo vệ không bị đổi giá trị. Nó không chứng minh mask bảo vệ hoàn hảo về mặt lâm sàng. "
        "Vì vậy review panel vẫn cần được kiểm tra, đặc biệt ở dị vật chồng trực tiếp lên xương.",
        ORANGE,
    )

    add_heading(doc, "6. Mô hình tuổi xương dùng để so sánh", 1)
    add_heading(doc, "6.1. Mô hình được chọn", 2)
    add_text(
        doc,
        "Mô hình đánh giá là checkpoint mở ianpan/bone-age trên Hugging Face [2], được cố định ở revision 2ab81275b84e9f518f04584177221a1d8c1dc1a5. "
        "Mô hình có khoảng 84,1 triệu tham số và là ensemble ba fold dựa trên ConvNeXt V2 Tiny. Đầu vào gồm ảnh X-quang thang xám và một kênh mã hóa giới tính.",
    )
    add_text(
        doc,
        "Lý do lựa chọn: (1) trọng số và mã inference công khai; (2) hỗ trợ đúng giới tính; (3) model card công bố kết quả trên chính 200 ca RSNA; "
        "(4) có bốn protocol tiền xử lý để kiểm tra độ bền; (5) có thể khóa revision và tái lập. Việc chạy lại của nhóm thu MAE ảnh gốc 4,4199 tháng, khớp con số 4,42 trong model card.",
    )
    add_callout(
        doc,
        "Phân biệt với mô hình trong bài báo chính",
        "Matsuoka và cộng sự dùng ensemble ResNet50 phân theo giới tính và báo cáo MAE gốc 6,26 tháng. "
        "ianpan/bone-age là checkpoint ConvNeXt V2 khác. Vì vậy không được đặt 30,11 tháng của bài báo và 43,40 tháng của thí nghiệm này như hai phép đo của cùng một mô hình; chỉ có thể so sánh xu hướng suy giảm.",
        RED,
    )

    add_heading(doc, "6.2. Mô hình đã được đề cập trong bài báo nào?", 2)
    add_text(
        doc,
        "ConvNeXt V2 được giới thiệu bởi Woo và cộng sự trong bài “ConvNeXt V2: Co-designing and Scaling ConvNets with Masked Autoencoders” [3]. "
        "Checkpoint ianpan/bone-age hiện được phát hành chủ yếu qua model card có DOI riêng trên Hugging Face; chưa tìm thấy một bài báo bình duyệt mô tả chính xác checkpoint ConvNeXt V2 ba fold này. "
        "Tác giả Ian Pan có công trình trước về ensemble tuổi xương trong RSNA AI [4], nhưng không nên đồng nhất công trình đó với checkpoint hiện tại.",
    )

    add_heading(doc, "6.3. Cách mô hình hoạt động", 2)
    add_number(doc, "Ảnh được co theo cạnh dài tối đa 512 pixel, sau đó padding thành 512 × 512; không bóp méo tỷ lệ.")
    add_number(doc, "Kênh 1 là ảnh thang xám; kênh 2 là giới tính (female = 255, male = 0), rồi chuẩn hóa về khoảng gần [−1, 1].")
    add_number(doc, "ConvNeXt V2 Tiny trích xuất feature map 768 chiều. Kiến trúc dùng các block convolution hiện đại và Global Response Normalization để tăng tương tác kênh.")
    add_number(doc, "GeM pooling với p = 3 tổng hợp không gian; so với average pooling, GeM nhấn mạnh vùng đáp ứng mạnh nhưng vẫn khả vi.")
    add_number(doc, "Dropout 0,1 và lớp tuyến tính tạo 240 logit, tương ứng 240 lớp tuổi theo tháng.")
    add_number(doc, "Softmax biến logit thành xác suất pᵢ; tuổi dự đoán là kỳ vọng Σpᵢ·i. Ba fold được lấy trung bình để giảm phương sai.")
    add_equation(doc, "ŷ = (1/3) Σₖ [ Σᵢ softmax(zₖ)ᵢ · i ] ,  i = 0,…,239")
    add_text(
        doc,
        "Biểu diễn bài toán như phân loại có thứ tự rồi lấy kỳ vọng giúp mô hình tạo dự đoán liên tục theo tháng, đồng thời phân bố softmax phản ánh mức tin cậy tương đối ở các tuổi lân cận.",
    )

    add_heading(doc, "7. Thiết kế thực nghiệm và ý nghĩa các phép đo", 1)
    add_heading(doc, "7.1. Ba nhóm ảnh và bốn protocol", 2)
    add_text(
        doc,
        "Mọi phép so sánh đều ghép cặp theo Case_ID và dùng cùng giới tính. Ba nhóm là: ảnh gốc; ảnh làm sạch artifact-only; ảnh tác giả, trong đó dự đoán của ba lần sinh được lấy trung bình theo ca. "
        "Bốn protocol giúp kiểm tra liệu kết luận có phụ thuộc vào histogram hay crop hay không.",
    )
    add_table(
        doc,
        ["Protocol", "Thao tác", "Câu hỏi kiểm tra"],
        [
            ["Toàn khung", "Resize/pad trực tiếp", "Tác động tổng thể trên ảnh nguyên bố cục?"],
            ["Histogram", "Khớp phân bố cường độ trước inference", "Sai khác có chỉ do độ sáng/tương phản?"],
            ["Crop cố định", "Crop theo cùng mask bảo vệ + lề 3%", "Sai khác có do nền ngoài bàn tay?"],
            ["Crop + histogram", "Kết hợp hai thao tác", "Khi giảm cả nhiễu bố cục và cường độ, kết luận còn giữ?"],
        ],
        [3.2, 6.3, 6.9],
        8.5,
    )

    add_heading(doc, "7.2. Các chỉ số", 2)
    add_text(doc, "Sai số tuyệt đối trung bình (MAE) đo độ lệch trung bình theo tháng, ít bị chi phối bởi một vài ngoại lệ hơn RMSE:")
    add_equation(doc, "MAE = (1/N) Σ |ŷᵢ − yᵢ|")
    add_text(doc, "RMSE bình phương sai số trước khi lấy trung bình nên phạt mạnh các ca sai rất lớn:")
    add_equation(doc, "RMSE = √[(1/N) Σ (ŷᵢ − yᵢ)²]")
    add_text(doc, "Bias là sai số có dấu. Bias dương nghĩa là mô hình dự đoán già hơn ground truth; bias âm nghĩa là dự đoán trẻ hơn:")
    add_equation(doc, "Bias = (1/N) Σ (ŷᵢ − yᵢ)")
    add_text(
        doc,
        "Within-6/12/24 là tỷ lệ ca có |ŷ−y| không vượt quá 6, 12 hoặc 24 tháng. Pearson đo mức bảo toàn thứ tự tuyến tính giữa dự đoán và ground truth, nhưng Pearson cao không đồng nghĩa MAE thấp.",
    )
    add_text(
        doc,
        "Độ trôi dự đoán ghép cặp là |ŷ_processed−ŷ_original|. Đây là chỉ số trực tiếp nhất cho câu hỏi “tiền xử lý đã làm mô hình đổi quyết định bao nhiêu tháng?”. "
        "ΔAE = |ŷ_processed−y| − |ŷ_original−y|; giá trị âm là cải thiện, dương là xấu đi.",
    )
    add_equation(doc, "Drift = |ŷprocessed − ŷoriginal| ;   ΔAE = AEprocessed − AEoriginal")

    add_heading(doc, "7.3. Bootstrap và Wilcoxon", 2)
    add_text(
        doc,
        "Khoảng tin cậy 95% được tính bằng paired bootstrap 10.000 lần: lấy mẫu lại 200 Case_ID có hoàn lại, giữ nguyên cặp ảnh gốc–ảnh xử lý trong mỗi lần. "
        "Phân vị 2,5% và 97,5% của phân bố thống kê bootstrap tạo CI. Nếu CI của ΔAE chứa 0, chưa có bằng chứng chắc chắn rằng độ chính xác thay đổi.",
    )
    add_text(
        doc,
        "Wilcoxon signed-rank kiểm tra phân bố sai số tuyệt đối ghép cặp mà không giả định chuẩn. p nhỏ cho biết sai khác có hệ thống, nhưng không tự nói hiệu ứng lớn hay có ý nghĩa lâm sàng; cần đọc cùng độ trôi và CI.",
    )

    add_heading(doc, "7.4. Cách đọc các loại biểu đồ", 2)
    add_bullet(doc, "Scatter ground truth–prediction: trục x là tuổi thật (tháng), trục y là tuổi dự đoán. Đường y=x là dự đoán hoàn hảo; điểm càng xa đường càng sai.")
    add_bullet(doc, "Boxplot sai số tuyệt đối: trục y là |ŷ−y| theo tháng; đường giữa hộp là median, hộp là IQR, râu và điểm ngoài cho thấy độ phân tán/outlier.")
    add_bullet(doc, "Bland–Altman: trục x là trung bình hai dự đoán (gốc và xử lý), trục y là processed−original. Đường giữa là bias; hai đường ngoài là bias ±1,96 SD, gọi là giới hạn đồng thuận.")
    add_bullet(doc, "Biểu đồ MAE: cột thấp hơn tốt hơn; thanh lỗi là CI 95% bootstrap. Cần so trong cùng protocol, không chỉ so chiều cao tuyệt đối giữa protocol.")
    add_bullet(doc, "Biểu đồ variability của tác giả: mỗi ca có độ lệch chuẩn giữa ba lần sinh. Trục y lớn cho biết prompt/mask giống nhau nhưng kết quả mô hình tuổi xương không ổn định.")

    add_heading(doc, "8. Kết quả so sánh định lượng và phân tích biểu đồ", 1)
    summary = load_summary()
    paired = load_paired()
    table_rows = []
    for r in summary:
        table_rows.append([
            protocol_label(r["Protocol"]),
            group_label(r["Group"]),
            f'{float(r["MAE_Months"]):.2f}',
            f'{float(r["RMSE_Months"]):.2f}',
            f'{float(r["Bias_Months"]):+.2f}',
            f'{100*float(r["Within_12_Months"]):.1f}%',
            f'{float(r["Pearson_vs_Truth"]):.3f}',
        ])
    add_table(
        doc,
        ["Protocol", "Nhóm", "MAE", "RMSE", "Bias", "≤12 tháng", "r"],
        table_rows,
        [3.7, 3.4, 1.7, 1.7, 1.7, 2.2, 1.8],
        7.9,
    )
    add_figure(
        doc,
        ASSET / "mae_three_way.png",
        "Hình 5. MAE của ba nhóm trên bốn protocol. Trục x là protocol; trục y là MAE theo tháng. Thanh lỗi là CI 95% bootstrap. Ảnh gốc và ảnh làm sạch gần như chồng nhau, trong khi ảnh tác giả cao hơn khoảng mười lần.",
        16.2,
    )
    add_text(
        doc,
        "Nhận xét tổng quát: ảnh làm sạch không làm suy giảm MAE trong bất kỳ protocol nào; MAE còn thấp hơn ảnh gốc từ 0,10 đến 0,20 tháng. "
        "Mức chênh nhỏ này không nên được diễn giải là phương pháp chắc chắn “cải thiện mô hình”, vì mục tiêu chính là bảo toàn; chỉ protocol crop+histogram cho Wilcoxon p=0,004 và CI ΔAE không chứa 0. "
        "Ngược lại, ảnh tác giả gây mức tăng MAE 38,5–40,6 tháng, nhất quán ở cả bốn protocol.",
    )
    add_figure(
        doc,
        ASSET / "paired_drift_three_way.png",
        "Hình 6. Độ trôi dự đoán so với ảnh gốc. Trục y là trung bình |dự đoán xử lý − dự đoán gốc| theo tháng. Ảnh làm sạch chỉ trôi 0,51–0,76 tháng; ảnh tác giả trôi 41,69–43,34 tháng.",
        16.2,
    )

    protocol_notes = {
        "fullframe": (
            "Ảnh gốc đạt MAE 4,42 tháng; ảnh làm sạch 4,32 tháng; ảnh tác giả 43,40 tháng. "
            "Độ trôi của chúng ta 0,76 tháng (CI 0,64–0,89), Pearson với dự đoán gốc 0,9998 và ΔAE −0,10 tháng có CI chứa 0. "
            "Kết luận: toàn khung cho thấy thao tác làm sạch hầu như không đổi quyết định mô hình, còn inpainting tác giả tạo bias già hóa +43,00 tháng.",
        ),
        "histmatch": (
            "Sau khớp histogram, MAE ảnh tác giả vẫn 42,89 tháng và bias +42,37 tháng. "
            "Điều này bác bỏ giả thuyết rằng suy giảm chỉ do sáng/tối toàn cục. Ảnh làm sạch có drift 0,56 tháng và ΔAE −0,10 tháng, CI vẫn chứa 0.",
        ),
        "fixed_crop": (
            "Crop cố định loại phần lớn nền ngoài bàn tay nhưng ảnh tác giả vẫn có MAE 44,72 tháng. "
            "Vì vậy sai khác không chỉ đến từ nhãn hoặc bố cục nền; nó nằm trong nội dung bàn tay đã được sinh lại. Ảnh làm sạch có drift 0,68 tháng và p Wilcoxon 0,28.",
        ),
        "fixed_crop_histmatch": (
            "Đây là protocol mạnh nhất đối với checkpoint: ảnh gốc MAE 4,19 và ảnh làm sạch 4,00 tháng. "
            "ΔAE của chúng ta −0,20 tháng, CI −0,32 đến −0,08 và p=0,004, nhưng kích thước hiệu ứng rất nhỏ. "
            "Ảnh tác giả vẫn MAE 44,78 tháng, chứng tỏ chuẩn hóa crop và histogram không phục hồi được giải phẫu bị thay đổi.",
        ),
    }
    fig_no = 7
    for pcode in ["fullframe", "histmatch", "fixed_crop", "fixed_crop_histmatch"]:
        add_heading(doc, f"8.{['fullframe','histmatch','fixed_crop','fixed_crop_histmatch'].index(pcode)+1}. {protocol_label(pcode)}", 2)
        add_figure(
            doc,
            ASSET / f"{pcode}_three_way.png",
            f"Hình {fig_no}. Phân tích ba nhóm ở protocol {protocol_label(pcode).lower()}. Panel trái là scatter tuổi thật–dự đoán; panel phải là boxplot sai số tuyệt đối. Đường chéo ở scatter biểu diễn dự đoán lý tưởng.",
            16.2,
        )
        fig_no += 1
        add_text(doc, protocol_notes[pcode])
        add_figure(
            doc,
            ASSET / f"{pcode}_bland_altman_three_way.png",
            f"Hình {fig_no}. Bland–Altman ở protocol {protocol_label(pcode).lower()}. Trục x là trung bình dự đoán gốc và xử lý; trục y là xử lý trừ gốc. Cụm điểm của chúng ta quanh 0 thể hiện đồng thuận; cụm ảnh tác giả dịch mạnh lên phía dương thể hiện dự đoán già hơn.",
            16.2,
        )
        fig_no += 1

    add_heading(doc, "8.5. Biến thiên giữa ba lần sinh của tác giả", 2)
    add_figure(
        doc,
        ASSET / "author_variability.png",
        f"Hình {fig_no}. Độ biến thiên dự đoán giữa ba ảnh sinh của từng ca ở protocol toàn khung. Mỗi điểm/cột đại diện một Case_ID; trục y là độ lệch chuẩn tuổi dự đoán theo tháng. Trung bình SD là 13,33 tháng, lớn hơn con số 8,77 tháng trong bài báo vì checkpoint đánh giá khác.",
        16.2,
    )
    fig_no += 1
    add_text(
        doc,
        "Độ lệch chuẩn trung bình 13,33 tháng, median 9,79 tháng và cực đại 70,70 tháng. "
        "Điều này cho thấy cùng một ảnh đầu vào và cùng yêu cầu loại nhãn có thể dẫn tới ba ảnh khiến mô hình tuổi xương phản ứng rất khác nhau. "
        "Sự bất định này là một vấn đề riêng, ngoài bias trung bình: ngay cả khi hiệu chỉnh bias, từng lần sinh vẫn không đáng tin cậy.",
    )

    add_heading(doc, "8.6. Đối chiếu với kết quả bài báo", 2)
    add_table(
        doc,
        ["Nguồn", "Mô hình đánh giá", "Ảnh gốc MAE", "Ảnh tác giả MAE", "Ghi chú"],
        [
            ["Matsuoka et al. [1]", "Ensemble ResNet50 theo giới tính", "6,26", "30,11", "Kết quả công bố; calibration 19,55; 600 ảnh sinh."],
            ["Thí nghiệm của nhóm", "ianpan ConvNeXt V2, revision cố định", "4,42", "43,40", "Ảnh tác giả lấy TB ba dự đoán/ca; fullframe."],
        ],
        [3.3, 4.7, 2.4, 2.8, 3.0],
        8.3,
    )
    add_text(
        doc,
        "Hai hàng không dùng để khẳng định mô hình nào tốt hơn vì kiến trúc, preprocessing và calibration khác nhau. "
        "Giá trị của phép đối chiếu là tính nhất quán định tính: cả hai mô hình độc lập đều cho thấy inpainting toàn ảnh làm suy giảm nghiêm trọng ước lượng tuổi xương.",
    )

    add_heading(doc, "9. Cách cài đặt, chạy và tái lập kết quả", 1)
    add_heading(doc, "9.1. Cấu trúc đầu ra quan trọng", 2)
    add_table(
        doc,
        ["Thành phần", "Đường dẫn tương đối", "Nội dung"],
        [
            ["Pipeline", "artifact_only_pipeline/", "Detector, protection, reconstruction, QC."],
            ["Chạy làm sạch", "run_artifact_only.py", "Sinh đủ ảnh, mask, panel và CSV."],
            ["Annotation", "annotation_server.py + annotation_ui/", "Duyệt/vẽ mask dị vật."],
            ["Mask thủ công", "annotations/manual_artifacts/", "185 mask đã duyệt."],
            ["Ảnh cuối", "outputs/artifact_only_200_manual_v3/", "200 ảnh sạch và dữ liệu QC."],
            ["Benchmark", "run_open_boneage_benchmark.py", "Inference 3 nhóm × 4 protocol."],
            ["Kết quả benchmark", "outputs/boneage_benchmark_all_200/", "Prediction, summary, bootstrap."],
            ["Biểu đồ báo cáo", "outputs/technical_report_assets/", "Ảnh biểu đồ và CSV ba nhánh."],
        ],
        [3.3, 6.2, 6.3],
        8.2,
    )

    add_heading(doc, "9.2. Khởi động công cụ vẽ mask", 2)
    add_callout(
        doc,
        "Lệnh PowerShell",
        "cd C:\\Users\\ASUS\\Documents\\tao_sinh_anh_x_ray_hand\n"
        ".\\.venv\\Scripts\\python.exe .\\annotation_server.py\n"
        "Mở trình duyệt tại: http://127.0.0.1:8765/",
        PALE_BLUE,
    )
    add_text(
        doc,
        "Sau khi khởi động lại máy phải chạy lại server. Không mở trực tiếp annotation_ui/index.html bằng file://. "
        "Trong giao diện, chọn cọ để tô đúng vùng nhãn/marker, dùng tẩy để sửa, “Lưu mask” cho ca có dị vật hoặc “Duyệt không sửa” nếu không cần mask.",
    )

    add_heading(doc, "9.3. Chạy pipeline làm sạch", 2)
    add_callout(
        doc,
        "Lệnh tham khảo",
        "python run_artifact_only.py --manual-mask-dir annotations\\manual_artifacts "
        "--output-dir outputs\\artifact_only_200_manual_v3\n"
        "Sau khi chạy, kiểm tra file summary/QC và thư mục review trước khi benchmark.",
        PALE_BLUE,
    )
    add_text(
        doc,
        "Tên tham số cụ thể có thể xem bằng python run_artifact_only.py --help. Không ghi đè dữ liệu gốc; mỗi lần thử nên dùng một output-dir mới để giữ lịch sử thí nghiệm.",
    )

    add_heading(doc, "9.4. Cài và chạy mô hình tuổi xương", 2)
    add_callout(
        doc,
        "Môi trường đề nghị",
        "python -m venv .boneage_env\n"
        ".\\.boneage_env\\Scripts\\Activate.ps1\n"
        "python -m pip install -r requirements_boneage.txt\n"
        "Checkpoint: ianpan/bone-age, revision 2ab81275b84e9f518f04584177221a1d8c1dc1a5",
        PALE_BLUE,
    )
    add_text(
        doc,
        "Máy cần PyTorch tương thích CUDA nếu chạy GPU; CPU vẫn có thể chạy nhưng chậm. Lần đầu Hugging Face phải tải trọng số và cần mạng. "
        "Manifest của thí nghiệm đầy đủ dùng torch 2.7.1+cu128, CUDA và bootstrap 10.000 lần.",
    )
    add_callout(
        doc,
        "Benchmark ba nguồn ảnh",
        "python run_open_boneage_benchmark.py --help\n"
        "Thiết lập nhóm original, ours và ba template author_1/author_2/author_3; chạy bốn protocol "
        "fullframe, histmatch, fixed_crop, fixed_crop_histmatch; cố định cùng CSV tuổi/giới tính.",
        PALE_BLUE,
    )
    add_text(
        doc,
        "Một lần chạy đầy đủ gồm 200 ca × 5 nhóm ảnh × 4 protocol = 4.000 inference. "
        "Script tổng hợp sau đó lấy trung bình ba dự đoán tác giả theo từng ca, không trộn ảnh hay lấy trung bình pixel.",
    )

    add_heading(doc, "10. Hạn chế, khuyến nghị và kết luận cuối cùng", 1)
    add_heading(doc, "10.1. Hạn chế", 2)
    add_bullet(doc, "Mask thủ công tốn công và phụ thuộc người gán nhãn; nên đo thỏa thuận liên người gán trên một tập con.")
    add_bullet(doc, "Mask bảo vệ là heuristic, chưa phải segmentation lâm sàng được chuyên gia xác nhận.")
    add_bullet(doc, "Mô hình nền bậc một phù hợp nhãn ngoài vùng giải phẫu nhưng không thể khôi phục thông tin bị dị vật che khuất.")
    add_bullet(doc, "Kết quả tuổi xương dựa trên một checkpoint mở; cần kiểm tra thêm mô hình thứ hai nếu muốn kết luận tổng quát.")
    add_bullet(doc, "Bộ 200 ca là test subset; không dùng để điều chỉnh tham số quá mức hoặc tuyên bố hiệu quả lâm sàng.")

    add_heading(doc, "10.2. Khuyến nghị bước tiếp theo", 2)
    add_number(doc, "Khóa phiên bản 200 mask và tạo file manifest gồm SHA-256 của ảnh gốc, mask và ảnh sạch.")
    add_number(doc, "Lấy mẫu ngẫu nhiên 30–50 ca để hai người độc lập review vùng dị vật và vùng bảo vệ.")
    add_number(doc, "Báo cáo cả kết quả chính toàn bộ 200 ca và phân tích nhạy cảm loại các ca dị vật chồng giải phẫu.")
    add_number(doc, "Chạy thêm một mô hình tuổi xương độc lập nếu có checkpoint hợp lệ; không chọn mô hình dựa trên kết quả thuận lợi.")
    add_number(doc, "Giữ phương án Stable Diffusion/ControlNet như ablation riêng, với invariant ngoài mask và đánh giá hallucination, không thay thế phương pháp chính.")

    add_heading(doc, "10.3. Nhận xét cuối cùng", 2)
    add_text(
        doc,
        "Quá trình phát triển cho thấy sai lầm ban đầu không nằm ở một kernel morphology cụ thể mà ở phạm vi bài toán: tách toàn bộ bàn tay để xóa vài nhãn nhỏ khiến lỗi phân đoạn bị phóng đại thành lỗi giải phẫu. "
        "Giải pháp artifact-only đảo ngược ưu tiên: giữ nguyên ảnh mặc định, chỉ cho phép thay đổi ở nơi đã chứng minh là dị vật.",
    )
    add_text(
        doc,
        "Kết quả 200/200 ảnh, bất biến pixel 200/200 và độ trôi tuổi xương dưới một tháng trên cả bốn protocol là bằng chứng nhất quán rằng pipeline đã đạt mục tiêu kỹ thuật. "
        "Trong khi đó, ảnh inpainting toàn ảnh của tác giả vẫn gây sai lệch lớn sau crop và khớp histogram, cho thấy vấn đề chủ yếu là nội dung giải phẫu bị tái tổng hợp chứ không đơn thuần là nền.",
    )
    add_callout(
        doc,
        "Kết luận",
        "Đối với tiền xử lý ảnh X-quang phục vụ ước lượng tuổi xương, chỉnh sửa cục bộ có mask bảo vệ, kiểm thử bất biến và review theo ca là hướng phù hợp hơn sinh lại toàn ảnh. "
        "Đây là đóng góp rõ ràng, có thể giải thích, có thể tái lập và khác về bản chất so với pipeline Stable Diffusion/ControlNet.",
        GREEN,
    )

    add_heading(doc, "Tài liệu tham khảo", 1)
    refs = [
        ("[1] Matsuoka F. và cộng sự. Evaluating the Clinical Impact of Generative Inpainting on Bone Age Estimation. arXiv:2511.23066, 2025.", "https://arxiv.org/abs/2511.23066"),
        ("[2] Ian Pan. ianpan/bone-age model card, Hugging Face, DOI 10.57967/hf/4673.", "https://huggingface.co/ianpan/bone-age"),
        ("[3] Woo S. và cộng sự. ConvNeXt V2: Co-designing and Scaling ConvNets with Masked Autoencoders. arXiv:2301.00808, 2023.", "https://arxiv.org/abs/2301.00808"),
        ("[4] Pan I. và cộng sự. Bone Age Assessment Using Deep Neural Networks with Synchronous and Asynchronous Learning. Radiology: AI, 2019.", "https://doi.org/10.1148/ryai.2019190053"),
        ("[5] Halabi S.S. và cộng sự. The RSNA Pediatric Bone Age Machine Learning Challenge. Radiology, 2019.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC6358027/"),
        ("[6] Bộ dữ liệu Synthetic Hand X-ray Dataset for Bone Age của Felipe Matsuoka trên Kaggle.", "https://www.kaggle.com/datasets/felipematsuoka/synthetic-hand-x-ray-dataset-for-bone-age/"),
        ("[7] Repo tham chiếu chính thức của nghiên cứu Matsuoka.", "https://github.com/felipe-matsuoka123/EVALUATING-THE-CLINICAL-IMPACT-OF-GENERATIVE-INPAINTING-ON-BONE-AGE-ESTIMATION"),
        ("[8] Fork chứa data/rsna_test.csv được dùng để đối chiếu ground truth.", "https://github.com/sRassmann/bone-age/blob/master/data/rsna_test.csv"),
        ("[9] Repo làm việc của nhóm gpt-image-bone-age-synthesis.", "https://github.com/nhattoan235/gpt-image-bone-age-synthesis"),
    ]
    for txt, url in refs:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.add_run(txt + " ")
        add_hyperlink(p, "Liên kết", url)

    add_heading(doc, "Phụ lục A. Quy tắc diễn giải và sử dụng kết quả", 1)
    add_bullet(doc, "Không gọi ảnh đầu ra là “ground truth giải phẫu”; đây là ảnh tiền xử lý bảo toàn pixel trong vùng được bảo vệ.")
    add_bullet(doc, "Không dùng MAE thấp hơn 0,1–0,2 tháng để tuyên bố lợi ích lâm sàng nếu chưa có external validation.")
    add_bullet(doc, "Không so trực tiếp MAE 30,11 và 43,40 như cùng một benchmark; hai mô hình đánh giá khác nhau.")
    add_bullet(doc, "Khi công bố, kèm manifest, revision model, protocol, seed/bootstrap và danh sách ca bị đánh dấu dị vật chồng giải phẫu.")

    add_heading(doc, "Phụ lục B. Checklist nghiệm thu 200 ảnh", 1)
    checklist = [
        "Đủ 200 Case_ID và tên đầu ra khớp ảnh gốc.",
        "Không còn nhãn/marker ngoài vùng giải phẫu theo review.",
        "Không mất đầu ngón, da, bờ cổ tay hoặc cẳng tay.",
        "Không có lỗ, halo, ghost chữ hoặc mảng nền phẳng bất thường.",
        "changed_inside_protected = 0.",
        "changed_outside_artifact = 0.",
        "Có panel Original / Mask / Cleaned / Difference cho từng ca.",
        "Benchmark dùng đúng giới tính, ground truth và revision checkpoint.",
    ]
    for item in checklist:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(2)
        p.add_run("☐ ").bold = True
        p.add_run(item)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
