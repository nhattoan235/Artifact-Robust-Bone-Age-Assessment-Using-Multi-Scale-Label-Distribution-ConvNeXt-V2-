from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"D:\Hoctap\Doan_totnghiep\NghienCuuChinh")
OUT = ROOT / "Bao_cao_do_an_tot_nghiep_du_thao.docx"
CONTENT_WIDTH_DXA = 8405  # leaves a 100-DXA table indent inside the A4 text area


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=10.5, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    if sum(widths) != CONTENT_WIDTH_DXA:
        scale = CONTENT_WIDTH_DXA / sum(widths)
        widths = [int(round(w * scale)) for w in widths]
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "100")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths[i]))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def set_font(run, size=13, bold=False, italic=False, color="000000"):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_page_number_format(section, fmt="decimal", start=1):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn("w:pgNumType"))
    if pgNumType is None:
        pgNumType = OxmlElement("w:pgNumType")
        sectPr.append(pgNumType)
    pgNumType.set(qn("w:fmt"), fmt)
    pgNumType.set(qn("w:start"), str(start))


def configure_section(section, page_fmt=None, include_page=True):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(0.2)
    section.footer_distance = Cm(0.8)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    for p in section.header.paragraphs:
        clear_paragraph(p)
    for p in section.footer.paragraphs:
        clear_paragraph(p)
    if include_page:
        add_page_number(section.footer.paragraphs[0])
        if page_fmt:
            set_page_number_format(section, page_fmt, 1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    set_font(run, 10)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    if level == 2:
        text = text.upper()
    r = p.add_run(text)
    set_font(r, {1: 18, 2: 14, 3: 13}[level], bold=True, color="000000")
    return p


def add_body(doc, text, bold_prefix=None, italic=False, align=None):
    p = doc.add_paragraph(style="Normal")
    p.alignment = align or WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.75)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_font(r1, 13, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_font(r2, 13, italic=italic)
    else:
        r = p.add_run(text)
        set_font(r, 13, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Cm(0.8 if level == 0 else 1.3)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    set_font(r, 13)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    set_font(r, 13)
    return p


def add_note(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.right_indent = Cm(0.15)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.1
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "FFF2CC")
    pPr.append(shd)
    pBdr = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "4")
        node.set(qn("w:color"), "D6B656")
        pBdr.append(node)
    pPr.append(pBdr)
    r = p.add_run("Lưu ý: " + text)
    set_font(r, 11)


def add_table(doc, headers, rows, widths, font_size=12):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    repeat_header(hdr)
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color="FFFFFF", size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(hdr.cells[i], "1F4E79")
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            if ri % 2 == 1:
                set_cell_shading(cells[i], "F3F6F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, 12, italic=False, color="000000")
    if text.upper().startswith("BẢNG"):
        previous = p._p.getprevious()
        if previous is not None and previous.tag == qn("w:p"):
            previous = previous.getprevious()
        if previous is not None and previous.tag == qn("w:tbl"):
            previous.addprevious(p._p)


def add_chapter_break(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run(title.upper())
    set_font(r, 18, bold=True, color="1F4E79")


def configure_document(doc):
    section = doc.sections[0]
    configure_section(section, include_page=False)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(6)
    for name, size, color, align in [("Heading 1", 18, "000000", WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", 14, "000000", WD_ALIGN_PARAGRAPH.LEFT), ("Heading 3", 13, "000000", WD_ALIGN_PARAGRAPH.LEFT)]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing = 1.3
        style.paragraph_format.space_after = Pt(6)


def add_cover_page(doc, draft_label=True):
    for _ in range(2):
        doc.add_paragraph()
    for text, size, bold in [
        ("BỘ CÔNG THƯƠNG", 13, True),
        ("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HCM\nKHOA CÔNG NGHỆ THÔNG TIN", 13, True),
        ("------", 12, False),
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); set_font(r, size, bold=bold)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ĐỒ ÁN KHÓA LUẬN TỐT NGHIỆP"); set_font(r, 18, bold=True)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DỰ ĐOÁN TUỔI XƯƠNG TỪ ẢNH X-QUANG BÀN TAY\nBẰNG MÔ HÌNH HỌC SÂU")
    set_font(r, 18, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Mã đề tài: CNTT-KLCN228"); set_font(r, 13, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Ngành: Công nghệ thông tin"); set_font(r, 13)
    for _ in range(2):
        doc.add_paragraph()
    for text in [
        "SINH VIÊN THỰC HIỆN:",
        "[MSSV – Họ tên – Lớp]",
        "[MSSV – Họ tên – Lớp]",
        "[MSSV – Họ tên – Lớp]",
        "GIẢNG VIÊN HƯỚNG DẪN: [CẦN BỔ SUNG]",
    ]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); set_font(r, 13, bold=text.endswith(":") or "SINH VIÊN" in text)
    if draft_label:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("BẢN DỰ THẢO – CẦN BỔ SUNG THÔNG TIN HÀNH CHÍNH"); set_font(r, 11, italic=True, color="666666")
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TP. HỒ CHÍ MINH, tháng 8 năm 2026"); set_font(r, 13)


def build():
    doc = Document()
    configure_document(doc)

    # Hai trang bìa theo form mẫu: bìa chính và bìa phụ.
    add_cover_page(doc, draft_label=True)
    front_cover = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front_cover, include_page=False)
    add_cover_page(doc, draft_label=False)
    front = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front, page_fmt="lowerRoman", include_page=True)

    # Executive front matter
    add_heading(doc, "LỜI CẢM ƠN", 1)
    add_body(doc, "Trong quá trình thực hiện đề tài, nhóm chúng em đã nhận được sự hướng dẫn của giảng viên, sự hỗ trợ của thầy cô, gia đình và bạn bè. Nhóm xin chân thành cảm ơn các thầy cô Khoa Công nghệ Thông tin, Trường Đại học Công Thương TP. HCM đã trang bị kiến thức nền tảng về lập trình, trí tuệ nhân tạo, xử lý ảnh và phương pháp nghiên cứu. Nhóm đặc biệt cảm ơn giảng viên hướng dẫn đã góp ý để nhóm xây dựng protocol thực nghiệm có kiểm soát, ghi nhận cả kết quả cải thiện và kết quả không đạt mục tiêu. Do thời gian và kinh nghiệm còn hạn chế, báo cáo khó tránh khỏi thiếu sót; nhóm kính mong nhận được nhận xét để tiếp tục hoàn thiện.")
    add_heading(doc, "NHẬN XÉT CỦA GIẢNG VIÊN HƯỚNG DẪN", 2)
    for _ in range(8):
        add_body(doc, "................................................................................................................................", align=WD_ALIGN_PARAGRAPH.LEFT)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("GIẢNG VIÊN HƯỚNG DẪN\n(Ký và ghi rõ họ tên)"); set_font(r, 13, bold=True)
    add_heading(doc, "TÓM TẮT", 1)
    add_body(doc, "Đề tài xây dựng và đánh giá một pipeline học sâu để dự đoán tuổi xương, tính theo tháng, từ ảnh X-quang bàn tay trẻ em. Nghiên cứu sử dụng bộ dữ liệu RSNA Pediatric Bone Age Challenge, xây dựng quy trình kiểm tra dữ liệu và chống rò rỉ, khảo sát augmentation, preprocessing, kiến trúc ConvNeXt và chiến lược ensemble 5-fold. Mô hình cuối được chọn bằng đánh giá out-of-fold trên 14.036 mẫu phát triển, không dùng nhãn test để chọn mô hình. Kết quả OOF đạt MAE 6,31669 tháng, RMSE 8,51942 và khoảng tin cậy bootstrap 95% [6,22464; 6,41133]. Một đánh giá thăm dò trên 200 ảnh test bằng ensemble đều cho MAE 4,73032 tháng; kết quả này được trình bày minh bạch là đánh giá sau khi đã mở nhãn test, không phải bằng chứng xác nhận mới.")
    add_body(doc, "Đóng góp chính của đồ án gồm: (i) khóa manifest và fingerprint cho train/validation/test; (ii) thiết kế trainer có checkpoint, resume, lưu RNG và audit; (iii) thực nghiệm có kiểm soát để chọn augmentation nhẹ A2 và độ phân giải 512; (iv) xây dựng báo cáo OOF, bootstrap confidence interval và phân tích theo giới tính/nhóm tuổi; (v) xác định các giới hạn cần giải quyết trong P9, đặc biệt là tái lập recipe preprocessing/training của các công trình gần đây.")
    add_heading(doc, "ABSTRACT", 1)
    add_body(doc, "This project develops and evaluates a deep-learning pipeline for pediatric bone-age estimation from hand radiographs. The RSNA Pediatric Bone Age Challenge data are used with an auditable split and leakage-control protocol. A pretrained ConvNeXt-Tiny backbone is conditioned on sex metadata and trained as a direct regression model with lightweight augmentation. Model selection is based on five-fold out-of-fold predictions over 14,036 development samples. The final OOF result is MAE 6.31669 months and RMSE 8.51942. A fixed equal-weight ensemble evaluated on 200 RSNA test images obtains an exploratory MAE of 4.73032 months. Because test labels were accessed for this exploratory evaluation, the result is not used for further model selection and is not claimed as a newly held-out confirmatory result.", italic=True)
    add_heading(doc, "DANH MỤC TỪ VIẾT TẮT", 1)
    add_table(doc, ["Viết tắt", "Giải thích"], [
        ["AI", "Artificial Intelligence – Trí tuệ nhân tạo"],
        ["CNN", "Convolutional Neural Network – Mạng nơ-ron tích chập"],
        ["OOF", "Out-of-fold prediction – Dự đoán ngoài fold huấn luyện"],
        ["MAE", "Mean Absolute Error – Sai số tuyệt đối trung bình"],
        ["RMSE", "Root Mean Squared Error – Căn bậc hai sai số bình phương trung bình"],
        ["CI", "Confidence Interval – Khoảng tin cậy"],
        ["GPU", "Graphics Processing Unit – Bộ xử lý đồ họa"],
    ], [1600, 7472], 12)
    add_heading(doc, "MỤC LỤC", 1)
    p = doc.add_paragraph("[[TOC]]")
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    # Chapter 1
    main = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(main, page_fmt="decimal", include_page=True)
    add_heading(doc, "MỞ ĐẦU", 1)
    add_body(doc, "Mở đầu trình bày lý do chọn đề tài, mục tiêu, nhiệm vụ, đối tượng, phạm vi và cấu trúc báo cáo. Đề tài tập trung vào việc xây dựng một pipeline học sâu có thể tái lập cho bài toán dự đoán tuổi xương từ ảnh X-quang bàn tay, sử dụng dữ liệu RSNA và đánh giá bằng các độ đo hồi quy phù hợp.")
    add_chapter_break(doc, "Chương 1. Tổng quan về đề tài")
    add_heading(doc, "1.1. Bối cảnh và lý do chọn đề tài", 2)
    add_body(doc, "Tuổi xương là một chỉ số phản ánh mức độ trưởng thành sinh học của trẻ, thường được ước lượng từ phim X-quang bàn tay bằng cách đối chiếu với atlas hoặc đánh giá của bác sĩ. Chỉ số này hỗ trợ đánh giá rối loạn tăng trưởng, theo dõi điều trị nội tiết và phân tích sự chênh lệch giữa tuổi xương với tuổi theo lịch. Tuy nhiên, đánh giá thủ công tốn thời gian, phụ thuộc kinh nghiệm và có thể có sai khác giữa các chuyên gia.")
    add_body(doc, "Sự phát triển của học sâu trong ảnh y tế tạo điều kiện xây dựng hệ thống hỗ trợ ước lượng tuổi xương tự động. Bài toán trên dữ liệu RSNA có ý nghĩa vì đây là benchmark công khai, có số lượng ảnh lớn và nhãn tuổi theo tháng. Đồ án lựa chọn hướng nghiên cứu có kiểm soát: thay vì chỉ tìm một điểm số tốt, nhóm xây dựng quy trình kiểm tra dữ liệu, chống leakage, lưu vết cấu hình và báo cáo bất định.")
    add_heading(doc, "1.2. Mục tiêu và nhiệm vụ nghiên cứu", 2)
    for t in [
        "Khảo sát bài toán dự đoán tuổi xương từ ảnh X-quang bàn tay và các công trình liên quan.",
        "Xây dựng pipeline học sâu dựa trên ConvNeXt-Tiny, có sử dụng thông tin giới tính và huấn luyện hồi quy tuổi theo tháng.",
        "Thiết kế protocol thực nghiệm có khóa split, kiểm tra duplicate/leakage, checkpoint/resume và đánh giá OOF bằng MAE, RMSE, accuracy theo ngưỡng và bootstrap CI.",
        "Phân tích ưu nhược điểm, giới hạn và đề xuất hướng phát triển để tiến tới một ứng dụng minh họa có khả năng tái lập.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "1.3. Đối tượng, phạm vi và câu hỏi nghiên cứu", 2)
    add_body(doc, "Đối tượng nghiên cứu là ảnh X-quang bàn tay trẻ em và mô hình học sâu dự đoán một giá trị liên tục là tuổi xương tính theo tháng. Phạm vi hiện tại tập trung vào dữ liệu RSNA, ảnh PNG, input vuông 512×512, backbone ConvNeXt-Tiny pretrained ImageNet-1K, augmentation nhẹ và huấn luyện trên GPU. Đồ án chưa nhằm thay thế chẩn đoán của bác sĩ, chưa đánh giá trên quần thể bệnh lý bên ngoài RSNA và chưa triển khai giao diện người dùng hoàn chỉnh.")
    for q in [
        "Augmentation nhẹ có giúp cải thiện sai số trên validation hay không?",
        "Mask nền bàn tay, biến thể kiến trúc và độ phân giải cao có tạo cải thiện ổn định không?",
        "Hiệu năng OOF của pipeline cuối là bao nhiêu và khác nhau thế nào theo giới tính/nhóm tuổi?",
        "Kết quả trên 200 ảnh test cần được diễn giải ra sao khi nhãn test đã được mở để đánh giá?",
    ]:
        add_bullet(doc, q)
    add_heading(doc, "1.4. Đóng góp của đồ án", 2)
    add_body(doc, "Đóng góp của đồ án nằm ở tính có kiểm soát và khả năng tái lập của quy trình: mỗi giai đoạn đều có artifact, manifest, hash, cấu hình và quyết định giữ/loại. Nhóm không chỉ báo cáo điểm số mà còn ghi nhận các nhánh không cải thiện như masking B1, ConvNeXtV2 FCMAE và multi-scale fusion; nhờ đó hạn chế việc chọn kết quả thuận lợi sau khi thử nghiệm.")
    add_heading(doc, "1.5. Bố cục báo cáo", 2)
    add_body(doc, "Báo cáo gồm sáu chương chính. Chương 1 trình bày bối cảnh và mục tiêu; Chương 2 trình bày cơ sở lý thuyết; Chương 3 mô tả dữ liệu, mô hình và protocol; Chương 4 trình bày cài đặt; Chương 5 báo cáo kết quả; Chương 6 thảo luận giới hạn và hướng phát triển. Cuối báo cáo là kết luận, kế hoạch 12 tuần, phân công nhóm, tài liệu tham khảo và phụ lục artifact.")

    # Chapter 2
    add_chapter_break(doc, "Chương 2. Cơ sở lý thuyết và công trình liên quan")
    add_heading(doc, "2.1. Khái niệm tuổi xương", 2)
    add_body(doc, "Tuổi xương là ước lượng mức độ trưởng thành của hệ xương, có thể khác với tuổi theo lịch. Trong dữ liệu RSNA, nhãn được biểu diễn theo tháng. Bài toán trong đồ án được mô hình hóa là hồi quy: với ảnh X-quang x và giới tính s, mô hình fθ dự đoán tuổi y-hat theo công thức y-hat = fθ(x, s).")
    add_heading(doc, "2.2. Bộ dữ liệu RSNA Pediatric Bone Age Challenge", 2)
    add_body(doc, "Theo hồ sơ audit của dự án, split được khóa gồm 12.611 ảnh train, 1.425 ảnh validation chính thức và 200 ảnh test. Từ train và validation, nhóm sử dụng development pool 14.036 mẫu để chạy 5-fold OOF. Mỗi manifest lưu image ID, tuổi theo tháng, giới tính, đường dẫn ảnh, SHA-256 và trạng thái đọc ảnh.")
    add_table(doc, ["Split", "Số ảnh", "Vai trò", "Quy tắc sử dụng"], [
        ["Train", "12.611", "Fit trọng số", "Được dùng trong huấn luyện"],
        ["Validation", "1.425", "So sánh ứng viên", "Dùng cho phát triển/đánh giá paired"],
        ["Development 5-fold", "14.036", "OOF cuối", "Mỗi mẫu có đúng một dự đoán ngoài fold"],
        ["RSNA test", "200", "Đánh giá thăm dò", "Không dùng chọn mô hình; đã mở nhãn trong P8"],
    ], [1800, 1300, 2300, 3672])
    add_caption(doc, "Bảng 2.1. Các split và vai trò trong protocol của dự án.")
    add_heading(doc, "2.3. Hồi quy ảnh với mạng nơ-ron tích chập", 2)
    add_body(doc, "CNN học các bộ lọc không gian để trích xuất cạnh, cấu trúc và đặc trưng hình ảnh. Với bài toán hồi quy, lớp cuối trả về một số thực thay vì xác suất lớp. Trong pipeline, tuổi được chuẩn hóa theo trung bình và độ lệch chuẩn của train; hàm mất mát Smooth L1 giúp giảm ảnh hưởng của các mẫu có sai số lớn trong giai đoạn tối ưu.")
    add_heading(doc, "2.4. ConvNeXt-Tiny và thông tin giới tính", 2)
    add_body(doc, "ConvNeXt là một họ ConvNet hiện đại được thiết kế lại theo các thực hành huấn luyện của mô hình thị giác giai đoạn 2020s. Đồ án sử dụng ConvNeXt-Tiny pretrained ImageNet-1K làm backbone. Đặc trưng ảnh được nối với embedding tuyến tính của giới tính, sau đó đưa qua một MLP gồm lớp ẩn, GELU, dropout và lớp hồi quy một chiều.")
    add_heading(doc, "2.5. Augmentation và preprocessing", 2)
    add_body(doc, "Augmentation A2 gồm lật ngang với xác suất 0,5, xoay trong khoảng ±7°, tịnh tiến tối đa 3%, scale 0,95–1,05, thay đổi brightness/contrast ±10% và gamma 0,9–1,1. B1 là nhánh mask nền bàn tay bằng mô hình segmentation, được đánh giá riêng và loại khỏi pipeline chính vì không cải thiện validation ổn định. Như vậy, preprocessing=none được giữ cho P7 final.")
    add_heading(doc, "2.6. Độ đo đánh giá", 2)
    add_body(doc, "Với N ảnh, MAE = (1/N) Σ|yᵢ − y-hatᵢ| và RMSE = sqrt((1/N) Σ(yᵢ − y-hatᵢ)²). Ngoài ra, nhóm báo cáo median absolute error và tỷ lệ dự đoán có sai số không vượt quá 6, 12 và 18 tháng. Khoảng tin cậy 95% của MAE được ước lượng bằng bootstrap 10.000 lần với seed 2026. Các so sánh paired trên cùng ảnh dùng delta absolute error; delta âm có lợi cho ứng viên.")
    add_heading(doc, "2.7. Công trình liên quan", 2)
    add_body(doc, "RSNA Pediatric Bone Age Challenge đặt nền tảng cho việc đánh giá tự động tuổi xương trên ảnh bàn tay. Các công trình sau đó báo cáo cải thiện bằng ensemble, mô hình học sâu và preprocessing. Trong hồ sơ dự án, hai mốc tham khảo trên cùng nominal test 200 ảnh là Deeplasia 3,87 tháng và công trình Bram năm 2025 là 3,68 tháng. Đây là các giá trị tham khảo theo literature, không phải phép so sánh hoàn toàn đồng nhất nếu khác preprocessing, checkpoint, cách lấy ground truth hoặc protocol.")

    # Chapter 3
    add_chapter_break(doc, "Chương 3. Phương pháp nghiên cứu và thiết kế thực nghiệm")
    add_heading(doc, "3.1. Nguyên tắc thiết kế protocol", 2)
    add_body(doc, "Protocol được thiết kế theo nguyên tắc khóa trước khi đánh giá: manifest, hash, split, cấu hình và tiêu chí giữ/loại phải được ghi lại. Test ground truth không được sử dụng để chọn augmentation, preprocessing, loss, trọng số ensemble hay hyperparameter. Sau khi P8 đã đọc nhãn test để đánh giá, mọi cải tiến tiếp theo phải chọn bằng validation/OOF; P8 chỉ được gọi là kết quả thăm dò.")
    add_heading(doc, "3.2. Kiểm tra dữ liệu và leakage", 2)
    for t in [
        "Kiểm tra số lượng ảnh, ảnh hỏng, ID trùng và SHA-256 trùng.",
        "Đối chiếu ID, tuổi và giới tính giữa manifest với annotation chính thức.",
        "Xác nhận không có ID/SHA giao nhau giữa train, validation và test.",
        "Lưu fingerprint và audit report để có thể kiểm tra lại trước khi chạy.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "3.3. Tiền xử lý và biểu diễn đầu vào", 2)
    add_body(doc, "Ảnh grayscale được pad thành hình vuông, resize về 512×512, lặp thành ba kênh để tương thích với backbone pretrained ImageNet, sau đó chuẩn hóa bằng mean/std ImageNet. Metadata giới tính được mã hóa M=1, F=0 và đưa qua embedding. Trong P7 final, ảnh được dùng ở trạng thái preprocessing=none; các nhánh segmentation/mask được đánh giá như ablation.")
    add_heading(doc, "3.4. Thuật toán và kiến trúc mô hình", 2)
    add_body(doc, "Gọi z = ConvNeXt(x) là vector đặc trưng ảnh và e = GELU(Wₛs + bₛ) là embedding giới tính. Vector điều kiện h = [z; e] được đưa qua MLP: u = GELU(W₁h + b₁), r = Dropout(u), y-hat_norm = W₂r + b₂. Giá trị tuổi theo tháng được khôi phục bằng y-hat = y-hat_norm × σ_train + μ_train. Cấu hình chính dùng sex_embedding_dim=16, hidden_dim=256, dropout=0,2.")
    add_heading(doc, "3.5. Huấn luyện và khả năng tái lập", 2)
    add_table(doc, ["Thành phần", "Cấu hình chính"], [
        ["Backbone", "ConvNeXt-Tiny pretrained ImageNet-1K"],
        ["Input", "512×512; grayscale lặp 3 kênh; ImageNet normalization"],
        ["Head", "Sex embedding 16 → MLP hidden 256 → regression"],
        ["Loss", "Smooth L1, beta = 3 tháng"],
        ["Optimizer/scheduler", "AdamW + cosine; learning rate 2e-4; weight decay 0,05"],
        ["Batch/effective batch", "Batch 12; gradient accumulation 3"],
        ["Huấn luyện", "Tối đa 35 epoch ở nhánh phát triển; early stopping patience 8"],
        ["Ổn định số", "Deterministic; AMP tự chọn BF16/FP16; gradient clipping 5"],
    ], [2600, 6472])
    add_caption(doc, "Bảng 3.1. Cấu hình nền tảng của pipeline ConvNeXt-Tiny.")
    add_heading(doc, "3.6. Chuỗi thực nghiệm", 2)
    add_table(doc, ["Phase", "Câu hỏi", "Quyết định"], [
        ["P0", "Dữ liệu có hợp lệ, không leakage?", "Khóa split và fingerprint; PASS"],
        ["P1", "Trainer/checkpoint có tái lập?", "Smoke/resume PASS; ưu tiên BF16 trên RTX 4050"],
        ["P2", "Augmentation nào phù hợp?", "Giữ A2: flip + biến đổi nhẹ"],
        ["P3", "Mask nền có giúp không?", "B1 không cải thiện; loại khỏi nhánh chính"],
        ["P4", "Kiến trúc mở rộng có tốt hơn?", "D0 đơn giản được giữ; D1/D2 loại; D3 phụ"],
        ["P5", "Kết quả có ổn định theo seed?", "D3 cải thiện nhỏ, CI chứa 0; chọn D0 đơn giản"],
        ["P6", "768 có đáng chi phí không?", "Chênh lệch -0,00137 tháng; giữ 512"],
        ["P7", "Kết quả OOF 5-fold?", "OOF 14.036 mẫu; MAE 6,31669 tháng"],
        ["P8", "Ensemble 5-fold trên test?", "MAE 4,73032 tháng; đánh giá thăm dò"],
    ], [1300, 3900, 3872])
    add_caption(doc, "Bảng 3.2. Lộ trình thực nghiệm và các quyết định giữ/loại.")
    add_heading(doc, "3.7. Đánh giá OOF và ensemble", 2)
    add_body(doc, "P7 dùng development pool 14.036 mẫu và chia 5 fold. Mỗi fold huấn luyện trên bốn phần và dự đoán phần còn lại; ghép lại tạo một dự đoán OOF duy nhất cho mỗi ID. P8 nạp best checkpoint của năm fold, suy luận trên 200 ảnh test và lấy trung bình đều năm dự đoán. Không có bước tối ưu trọng số ensemble trên test.")
    add_heading(doc, "3.8. Artifact và cấu trúc lưu trữ", 2)
    add_body(doc, "Mỗi run lưu config_resolved.yaml, environment.txt, metrics.jsonl, warnings.log, checkpoint, val_predictions_best.csv và run_state.json. Các file tổng hợp gồm P7_OOF_report.json, P7_5FOLD_AUDIT.txt, P8_test_ensemble_report.json và P8_ensemble_predictions.csv. Cấu trúc này hỗ trợ resume qua phiên GPU và cho phép kiểm tra tính nhất quán của kết quả.")

    # Chapter 4
    add_chapter_break(doc, "Chương 4. Cài đặt hệ thống")
    add_heading(doc, "4.1. Môi trường thực hiện", 2)
    add_table(doc, ["Hạng mục", "Mô tả"], [
        ["Ngôn ngữ", "Python 3.x"],
        ["Deep learning", "PyTorch, torchvision, timm"],
        ["Xử lý ảnh", "Pillow, OpenCV"],
        ["Tăng tốc", "CUDA GPU; kiểm tra smoke trên RTX 4050 và các phiên Colab"],
        ["Lưu trữ", "Workspace cục bộ, ZIP kết quả và persistent storage trong quá trình train"],
        ["Kiểm thử", "Preflight, smoke test, stop/resume, hash/audit và bootstrap evaluation"],
    ], [2600, 6472])
    add_heading(doc, "4.2. Các module chính", 2)
    add_table(doc, ["Module", "Chức năng"], [
        ["p0_audit", "Kiểm tra dataset, split, fingerprint và leakage"],
        ["p1_baseline/data.py", "Đọc manifest, pad/resize/normalize và augmentation"],
        ["p1_baseline/model.py", "ConvNeXt-Tiny, ConvNeXtV2, multi-scale và LDL head"],
        ["p1_baseline/train.py", "Huấn luyện, checkpoint, resume, metrics và warning"],
        ["p5_seed_confirmation", "Tổng hợp và so sánh theo seed"],
        ["p7_final_v3", "OOF 5-fold và audit cuối"],
        ["p8_test_ensemble", "Audit input và suy luận ensemble test"],
    ], [3000, 6072])
    add_heading(doc, "4.3. Luồng xử lý khi suy luận", 2)
    for t in [
        "Người dùng cung cấp ảnh X-quang và giới tính của ca cần dự đoán.",
        "Hệ thống chuyển ảnh grayscale, pad vuông, resize 512×512 và chuẩn hóa.",
        "Năm checkpoint fold tạo ra năm dự đoán tuổi theo tháng.",
        "Bộ suy luận lấy trung bình đều, ghi kết quả gồm image_id, sex, prediction_months và model_count.",
        "Kết quả cần được xem là hỗ trợ tham khảo, không phải chẩn đoán y khoa.",
    ]:
        add_number(doc, t)
    add_heading(doc, "4.4. Giao diện minh họa và trạng thái triển khai", 2)
    add_note(doc, "Hồ sơ hiện tại chưa chứa mã giao diện Streamlit/Gradio/Tkinter. Vì vậy phần này là đặc tả chức năng cần bổ sung, không được ghi nhận là sản phẩm đã hoàn thành.")
    add_table(doc, ["Chức năng dự kiến", "Đầu vào", "Đầu ra"], [
        ["Tải ảnh", "PNG/JPG X-quang bàn tay", "Ảnh xem trước và kiểm tra kích thước"],
        ["Nhập giới tính", "F/M", "Giá trị metadata được dùng bởi model"],
        ["Dự đoán", "Ảnh + giới tính + checkpoint ensemble", "Tuổi xương dự đoán theo tháng"],
        ["Lưu kết quả", "Mã ca và thời điểm", "CSV/JSON gồm prediction và cấu hình"],
        ["Cảnh báo", "Ảnh sai định dạng hoặc ngoài miền", "Thông báo lỗi, không trả kết quả giả"],
    ], [3000, 3000, 3072])
    add_heading(doc, "4.5. Yêu cầu an toàn và đạo đức", 2)
    add_body(doc, "Hệ thống không nên lưu thông tin định danh bệnh nhân nếu không cần thiết. Dữ liệu RSNA chỉ được sử dụng trong phạm vi học thuật và phi thương mại theo điều khoản dataset. Kết quả mô hình cần hiển thị nhãn ‘hỗ trợ nghiên cứu/tham khảo’, không diễn giải như chẩn đoán hoặc khuyến nghị điều trị.")

    # Chapter 5
    add_chapter_break(doc, "Chương 5. Kết quả thực nghiệm")
    add_heading(doc, "5.1. Kết quả audit dữ liệu", 2)
    add_body(doc, "P0 đạt PASS. Có 12.611 ảnh train, 1.425 ảnh validation, 200 ảnh test bị khóa nhãn trong giai đoạn phát triển và 14.036 mẫu development. Audit xác nhận ảnh đọc được, không có duplicate theo SHA-256 và không có ID/SHA giao nhau giữa các split. Validation chính thức được đối chiếu với annotation Deeplasia.")
    add_heading(doc, "5.2. Ảnh hưởng của augmentation", 2)
    add_table(doc, ["Ứng viên", "Mô tả", "Validation MAE (tháng)", "Quyết định"], [
        ["A0", "Không augmentation", "6,640 (xấp xỉ)", "Baseline"],
        ["A1", "Lật ngang", "6,514 (xấp xỉ)", "Không giữ riêng"],
        ["A2", "Lật + xoay/tịnh tiến/scale + brightness/contrast/gamma", "6,185", "Giữ làm pipeline chính"],
    ], [1100, 4000, 1800, 2172])
    add_body(doc, "A2 tốt hơn A0 khoảng 0,455 tháng trong so sánh paired; khoảng tin cậy bootstrap 95% của delta là [-0,658; -0,249]. Kết quả này là cơ sở để khóa augmentation nhẹ cho các phase tiếp theo.")
    add_heading(doc, "5.3. Preprocessing mask nền", 2)
    add_body(doc, "Nhánh B1 dùng mask toàn bàn tay bằng Efficient-UNet/TensorMask fallback. Validation MAE là 6,239 so với 6,185 của nhánh không mask; delta +0,054 tháng và khoảng tin cậy chứa 0. Theo quy tắc định trước, B1 không được đưa vào pipeline chính. Kết quả âm tính này quan trọng vì cho thấy thêm preprocessing không tự động tạo cải thiện.")
    add_heading(doc, "5.4. So sánh kiến trúc", 2)
    add_table(doc, ["Ứng viên", "Kết quả chính", "Diễn giải"], [
        ["D0 ConvNeXt-Tiny", "MAE khoảng 6,185 ở validation seed 42", "Baseline đơn giản, ổn định"],
        ["D1 ConvNeXtV2 FCMAE", "MAE khoảng 31,963", "Feature collapse trong recipe đã chạy; loại"],
        ["D2 multi-scale", "MAE khoảng 6,297", "Không cải thiện primary; loại"],
        ["D3 label-distribution", "MAE 6,145 seed 42", "Cải thiện nhỏ; CI/seed chưa đủ thuyết phục"],
    ], [2300, 2800, 3972])
    add_heading(doc, "5.5. Xác nhận seed và độ phân giải", 2)
    add_body(doc, "Trên ba seed 17/42/123, D0 có MAE trung bình 6,26229 và D3 fused có MAE trung bình 6,22508; delta trung bình -0,03721 tháng nhưng CI [-0,13521; 0,05863] chứa 0 và cải thiện dưới ngưỡng thực tiễn 0,10 tháng. Ở thí nghiệm độ phân giải, 768×768 đạt 6,18342 so với 6,18479 của 512×512; delta -0,00137 và CI [-0,19070; 0,19213]. Vì vậy nhóm chọn D0 ở 512×512 để giảm chi phí tính toán.")
    add_heading(doc, "5.6. Kết quả P7: 5-fold OOF", 2)
    add_table(doc, ["Chỉ số", "Kết quả"], [
        ["Số mẫu OOF / ID duy nhất", "14.036 / 14.036"],
        ["Pooled MAE", "6,316691 tháng"],
        ["RMSE", "8,519420"],
        ["Median absolute error", "4,875 tháng"],
        ["Accuracy ±6 / ±12 / ±18 tháng", "59,01% / 86,73% / 95,53%"],
        ["Bootstrap 95% CI của MAE", "[6,224638; 6,411327]"],
        ["MAE theo giới tính F / M", "6,53625 / 6,13108 tháng"],
        ["MAE nhóm tuổi 0–59 / 60–119 / 120–179 / 180–228", "6,06266 / 7,31682 / 5,91187 / 5,98890 tháng"],
    ], [4200, 4872])
    add_heading(doc, "5.7. Kết quả P8: ensemble trên 200 ảnh test", 2)
    add_body(doc, "P8 sử dụng trung bình đều dự đoán của năm fold, không tune trên test. Ensemble đạt MAE 4,730321 tháng, RMSE 6,028745, median absolute error 4,133363 tháng và accuracy ±6/±12/±18 lần lượt là 70,0%/94,0%/100,0%. Bootstrap 95% CI của MAE là [4,222475; 5,263709].")
    add_note(doc, "Diễn giải bắt buộc: nhãn test đã được đọc để tính metric P8. Vì vậy P8 là đánh giá thăm dò trên 200 ảnh, không phải một external hold-out mới và không được dùng để chọn mô hình tiếp theo.")
    add_heading(doc, "5.8. So sánh với mốc tham khảo", 2)
    add_table(doc, ["Nguồn tham khảo", "MAE báo cáo (tháng)", "So với P8", "Ghi chú"], [
        ["Bram et al. (2025)", "3,68", "P8 cao hơn 1,05", "Cần kiểm tra protocol đầy đủ trước khi so sánh kết luận"],
        ["Rassmann et al. – Deeplasia (2024)", "3,87", "P8 cao hơn 0,86", "Cùng nominal test 200; khác biệt có thể do recipe/ensemble"],
        ["Đồ án – P8 ensemble", "4,73032", "Mốc hiện tại", "Đánh giá thăm dò, CI rộng do n=200"],
    ], [2600, 1800, 1800, 2872])
    add_heading(doc, "5.9. Nhận xét về sai số", 2)
    add_body(doc, "OOF cho thấy nhóm tuổi 60–119 tháng có MAE cao nhất (7,31682 tháng), trong khi nhóm 120–179 tháng thấp hơn (5,91187 tháng). MAE ở nhóm F cao hơn M khoảng 0,405 tháng. Đây là mô tả phân tầng của kết quả, chưa đủ để kết luận có bias lâm sàng; cần thêm kích thước mẫu từng nhóm, calibration và đánh giá external để diễn giải chắc chắn hơn.")

    # Chapter 6
    add_chapter_break(doc, "Chương 6. Thảo luận, hạn chế và hướng phát triển")
    add_heading(doc, "6.1. Thảo luận kết quả", 2)
    add_body(doc, "Kết quả cho thấy augmentation nhẹ tạo cải thiện rõ ràng hơn so với các thay đổi kiến trúc phức tạp trong recipe hiện tại. Điều này phù hợp với đặc trưng bài toán: sai khác về hướng chụp, vị trí bàn tay và độ sáng có thể làm mô hình nhạy với biến thiên hình ảnh không liên quan trực tiếp tới tuổi xương. Ngược lại, mask nền và multi-scale không cải thiện, có thể do làm mất ngữ cảnh hoặc tăng độ khó tối ưu.")
    add_body(doc, "Việc chọn D0 thay vì D3 là quyết định dựa trên tính đơn giản và ổn định, không phải phủ nhận label-distribution learning nói chung. D3 chỉ được chạy theo một recipe, cải thiện trung bình nhỏ, CI qua 0 và không đạt ngưỡng 0,10 tháng. Đây là cách diễn giải phù hợp với dữ liệu thực nghiệm hiện có.")
    add_heading(doc, "6.2. Trả lời câu hỏi nghiên cứu", 2)
    for t in [
        "A2 có cải thiện: có, kết quả paired validation ủng hộ và được khóa.",
        "Mask/kiến trúc/độ phân giải cao có cải thiện ổn định: chưa có bằng chứng đủ mạnh trong các recipe đã chạy.",
        "Hiệu năng OOF: MAE 6,31669 tháng trên 14.036 mẫu với CI 95% [6,22464; 6,41133].",
        "Diễn giải test: P8 đạt 4,73032 tháng nhưng là đánh giá thăm dò sau khi đã mở nhãn test.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "6.3. Hạn chế", 2)
    for t in [
        "Chưa có external hold-out mới ngoài RSNA; P8 test 200 đã được mở nhãn.",
        "Hiệu năng OOF còn cao hơn mốc literature 3,68–3,87 tháng; không được tuyên bố vượt công trình.",
        "Chưa có giao diện người dùng hoàn chỉnh và chưa có đánh giá usability.",
        "Chưa có đánh giá bác sĩ, calibration, uncertainty ở mức từng ca hoặc kiểm tra trên bệnh lý xương bất thường.",
        "Kích thước test 200 làm khoảng tin cậy rộng; các so sánh literature cần đồng nhất protocol hơn.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "6.4. Hướng phát triển P9", 2)
    add_body(doc, "Hướng ưu tiên là tái lập công bằng recipe của Bram: preprocessing theo mô tả bài báo, augmentation và 100 epoch/early stopping; mọi cấu hình được chọn bằng validation/OOF. Nhánh P9-A1 mask Deeplasia đã early-stop với validation MAE 6,307675 và delta +0,122884 so với P2 A2; vì vậy chưa mở test và cần chạy A0 cùng recipe L1 để tách ảnh hưởng loss/preprocessing. Chỉ sau khi có OOF cải thiện lặp lại ở nhiều fold/seed mới cân nhắc mở một đánh giá test đã khóa trước.")
    add_heading(doc, "6.5. Tiêu chí nghiệm thu giai đoạn tiếp theo", 2)
    for t in [
        "Có cấu hình và manifest được hash; không dùng nhãn test để chọn.",
        "OOF MAE cải thiện rõ so với 6,31669 tháng và không có subgroup collapse.",
        "Cải thiện lặp lại ở ít nhất ba seed/fold hoặc có CI paired thuyết phục.",
        "Có giao diện suy luận tối thiểu và lưu được output, config, timestamp.",
    ]:
        add_bullet(doc, t)

    # Conclusion and planning
    add_chapter_break(doc, "Kết luận")
    add_body(doc, "Đồ án đã xây dựng được một pipeline nghiên cứu có kiểm soát cho bài toán dự đoán tuổi xương từ ảnh X-quang bàn tay. Nhóm đã khóa dữ liệu, kiểm tra leakage, khảo sát augmentation/preprocessing/kiến trúc/độ phân giải, xây trainer có resume và hoàn thành đánh giá OOF 5-fold. Mô hình ConvNeXt-Tiny với sex embedding, augmentation A2 và input 512 đạt MAE OOF 6,31669 tháng trên 14.036 mẫu. Ensemble 5-fold trên 200 ảnh test đạt 4,73032 tháng nhưng được phân loại đúng là đánh giá thăm dò.")
    add_body(doc, "Kết luận quan trọng nhất là tính tái lập và tính trung thực của protocol: các nhánh không hiệu quả vẫn được báo cáo, test không được dùng để chọn mô hình, và các mốc literature không bị dùng để tuyên bố vượt trội khi protocol chưa tương thích. Để hoàn thành sản phẩm theo đầy đủ rubric, nhóm cần bổ sung thông tin hành chính, giao diện minh họa, hình ảnh trực quan của pipeline và xác nhận của giảng viên hướng dẫn.")

    add_chapter_break(doc, "Kế hoạch thực hiện 12 tuần")
    schedule = [
        ["01", "Chốt đề tài, phạm vi, câu hỏi nghiên cứu và phân công nhóm", "Biên bản/outline được GVHD duyệt"],
        ["02", "Khảo sát RSNA, tuổi xương, metric và công trình liên quan", "Bảng tổng quan tài liệu"],
        ["03", "Audit dataset, split, duplicate, leakage và manifest", "P0 report PASS"],
        ["04", "Dựng baseline ConvNeXt-Tiny và trainer", "Preflight + smoke test"],
        ["05", "Kiểm thử checkpoint, resume, logging và reproducibility", "Recovery evidence"],
        ["06", "So sánh augmentation A0–A2", "Paired report và quyết định khóa A2"],
        ["07", "Khảo sát preprocessing/mask nền", "P3 report và phân tích âm tính"],
        ["08", "Khảo sát kiến trúc D0–D3", "P4 report"],
        ["09", "Seed confirmation và so sánh độ phân giải", "P5/P6 reports"],
        ["10", "Huấn luyện 5-fold, tạo OOF và audit", "P7 OOF report"],
        ["11", "Ensemble, phân tích kết quả, hoàn thiện giao diện minh họa", "P8 + prototype UI"],
        ["12", "Hoàn chỉnh báo cáo, slide, kiểm tra trích dẫn và nộp", "Bản cuối + slide + mã nguồn"],
    ]
    add_table(doc, ["Tuần", "Công việc", "Sản phẩm/tiêu chí"], schedule, [900, 5000, 3172])

    add_chapter_break(doc, "Phân công nhóm")
    add_note(doc, "Tên, MSSV và lớp chưa có trong thư mục dự án; nhóm cần thay các placeholder trước khi nộp.")
    add_table(doc, ["Thành viên", "Nhiệm vụ chính", "Sản phẩm chịu trách nhiệm"], [
        ["SV1 – [Họ tên/MSSV]", "Trưởng nhóm; protocol dữ liệu, leakage, quản lý artifact", "P0, P7, tổng hợp báo cáo"],
        ["SV2 – [Họ tên/MSSV]", "Mô hình, huấn luyện, checkpoint/resume, thực nghiệm kiến trúc", "P1, P4, P5, config/model"],
        ["SV3 – [Họ tên/MSSV]", "Tiền xử lý, đánh giá, trực quan hóa, giao diện minh họa", "P2, P3, P6, P8, prototype UI"],
    ], [2200, 4072, 2800])
    add_body(doc, "Cả ba thành viên cùng tích hợp kết quả vào một kho mã nguồn và một báo cáo chung trước mỗi buổi làm việc với giảng viên hướng dẫn. Phân công có thể điều chỉnh theo quyết định của GVHD.")

    add_chapter_break(doc, "Tài liệu tham khảo")
    refs = [
        "[1] Halabi SS, Prevedello LM, Kalpathy-Cramer J, et al., “The RSNA Pediatric Bone Age Machine Learning Challenge,” Radiology, 2018, 290(2):498–503.",
        "[2] Radiological Society of North America, “RSNA Pediatric Bone Age Challenge (2017),” RSNA AI Challenge, https://www.rsna.org/artificial-intelligence/ai-image-challenge/RSNA-Pediatric-Bone-Age-Challenge-2017.",
        "[3] Liu Z, Mao H, Wu CY, Feichtenhofer C, Darrell T, Xie S., “A ConvNet for the 2020s,” CVPR, 2022, doi:10.1109/CVPR52688.2022.01167.",
        "[4] Rassmann S, et al., “Deeplasia: deep learning for bone age assessment validated on skeletal dysplasias,” European Journal of Endocrinology, 2024. Giá trị RSNA test tham khảo trong hồ sơ dự án: MAE 3,87 tháng.",
        "[5] Bram et al., công trình năm 2025 về dự đoán tuổi xương trên RSNA; DOI được ghi trong hồ sơ dự án là 10.1177/03635465251353483. Cần nhóm kiểm tra lại metadata bài báo từ bản PDF trước khi nộp.",
        "[6] Kingma DP, Ba J., “Adam: A Method for Stochastic Optimization,” ICLR, 2015.",
        "[7] PyTorch Documentation, “PyTorch: An Open Source Machine Learning Framework,” https://pytorch.org/docs/.",
        "[8] timm documentation and model implementation, https://github.com/huggingface/pytorch-image-models.",
        "[9] Hồ sơ thực nghiệm nội bộ của nhóm: AI_Context/01_STATUS_RESULTS.md, 02_METHOD_HISTORY.md, 03_DATA_PROTOCOL.md, 04_NEXT_P9_PLAN.md và các artifact P0–P8 trong thư mục dự án.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(ref)
        set_font(r, 11)

    add_chapter_break(doc, "Phụ lục")
    add_heading(doc, "Phụ lục A. Bản đồ artifact quan trọng", 2)
    add_table(doc, ["Artifact", "Nội dung", "Trạng thái"], [
        ["p0_audit/outputs/P0_FINAL_REPORT.json", "Audit split và số lượng dữ liệu", "PASS"],
        ["p1_baseline/P2_FINAL_REPORT.json", "Quyết định augmentation A2", "PASS"],
        ["p5_seed_confirmation/P5_AGGREGATE.json", "So sánh seed D0/D3", "PASS"],
        ["p6_resolution/P6_768_vs_512_paired.json", "So sánh 768 với 512", "PASS"],
        ["p7_final_v3/P7_OOF_report.json", "5-fold OOF và CI", "PASS"],
        ["p8_test_ensemble/outputs/P8_test_ensemble_report.json", "Ensemble test 200 ảnh", "PASS – thăm dò"],
        ["p8_test_ensemble/outputs/P8_ensemble_predictions.csv", "Dự đoán từng ảnh test", "Artifact output"],
    ], [3500, 4072, 1600])
    add_heading(doc, "Phụ lục B. Danh sách thông tin cần nhóm bổ sung", 2)
    for t in [
        "Mã đề tài, tên trường/khoa/bộ môn và tên đề tài chính thức.",
        "Họ tên, MSSV, lớp của ba sinh viên.",
        "Họ tên, điện thoại, email giảng viên hướng dẫn.",
        "Ngày bắt đầu/kết thúc 12 tuần và lịch gặp GVHD.",
        "Ảnh chụp màn hình giao diện minh họa hoặc mã prototype sau khi triển khai.",
        "Thông tin GPU/phiên bản package chính xác của môi trường nộp bài.",
    ]:
        add_bullet(doc, t)
    add_note(doc, "Bản này là dự thảo báo cáo theo đúng các mục trong đề cương đúng mà người dùng cung cấp. Cần GVHD duyệt cách gọi ‘đồ án/khóa luận’, tài liệu tham khảo và phạm vi tuyên bố trước khi phát hành bản chính thức.")

    doc.core_properties.title = "Dự đoán tuổi xương từ ảnh X-quang bàn tay bằng mô hình học sâu"
    doc.core_properties.subject = "Báo cáo đồ án tốt nghiệp – bản dự thảo"
    doc.core_properties.author = "Nhóm sinh viên – cần bổ sung"
    doc.core_properties.comments = "Generated from project artifacts; administrative placeholders require completion."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
